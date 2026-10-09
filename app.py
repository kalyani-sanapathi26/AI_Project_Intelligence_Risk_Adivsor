import inspect
import streamlit as st
from src.document_loader import extract_text
from src.chunking import split_documents
from src.embeddings import generate_embeddings
from src.vector_store import store_embeddings
from src.rag_pipeline import retrieve_documents
from src.scope_agent import analyze_scope
from src.risk_agent import analyze_risks
from src.documentation_agent import (
    generate_project_documentation,
    _response_text,
    _fallback_user_stories,
    _format_risk_records,
    _clean_action_items,
    _format_project_summary,
)
from src.health_agent import analyze_project_health
from src.blocker_agent import analyze_blockers
from src.llm import generate_answer
import re
import html




st.set_page_config(
    page_title="AI Project Intelligence & Risk Advisor",
    page_icon="📊",
    layout="wide",
    initial_sidebar_state="expanded"
)

st.markdown("""
<style>
/* Interactive module cards */
.module-card {
    position: relative;
    overflow: hidden;
    border: 1px solid #38445c;
    border-left: 6px solid var(--accent, #7c9cff);
    border-radius: 18px;
    padding: 22px 24px;
    margin: 22px 0 10px 0;
    background: linear-gradient(135deg, #171d2b 0%, #20283a 100%);
    box-shadow: 0 8px 24px rgba(0,0,0,.18);
    transition: transform .22s ease, box-shadow .22s ease, border-color .22s ease;
}
.module-card:hover {
    transform: translateY(-4px);
    box-shadow: 0 14px 32px rgba(90,120,255,.20);
    border-color: #9cafff;
}
.module-card:before {
    content: "";
    position: absolute;
    width: 130px;
    height: 130px;
    right: -55px;
    top: -65px;
    border-radius: 50%;
    background: rgba(124,156,255,.08);
}
.module-icon { font-size: 34px; line-height: 1; margin-bottom: 10px; }
.module-title { color: #ffffff; font-size: 27px; font-weight: 800; margin: 0 0 8px 0; }
.module-description { color: #b2bdd0; font-size: 15px; line-height: 1.65; margin: 0 0 13px 0; }
.module-badge {
    display: inline-block;
    color: #c5d2ff;
    background: #263653;
    border: 1px solid #526da8;
    border-radius: 999px;
    padding: 5px 12px;
    font-size: 12px;
    font-weight: 700;
}
.module-line { height: 6px; background: #30394d; border-radius: 99px; margin-top: 16px; }
.module-line span { display:block; height:100%; width: var(--progress,70%); border-radius:99px; background: linear-gradient(90deg,var(--accent,#7c9cff),#b58cff); }
.module-caption { color:#8794aa; font-size:12px; margin-top:8px; }
div.stButton > button {
    border-radius: 12px;
    border: 1px solid #53617a;
    font-weight: 700;
    transition: all .2s ease;
}
div.stButton > button:hover {
    border-color: #9cafff;
    color: #ffffff;
    transform: translateY(-2px);
    box-shadow: 0 6px 16px rgba(124,156,255,.18);
}
</style>
""", unsafe_allow_html=True)



st.markdown("""
<style>

.section-card {
    background: linear-gradient(135deg, #171d2b, #20283a);
    border: 1px solid #39445c;
    border-left: 6px solid #7c9cff;
    border-radius: 18px;
    padding: 24px;
    margin: 18px 0;
    box-shadow: 0 8px 25px rgba(0, 0, 0, 0.20);
    transition: all 0.25s ease;
}

.section-card:hover {
    transform: translateY(-4px);
    border-color: #9cafff;
    box-shadow: 0 12px 32px rgba(90, 120, 255, 0.22);
}

.section-icon {
    font-size: 34px;
    margin-bottom: 8px;
}

.section-title {
    color: #ffffff;
    font-size: 25px;
    font-weight: 750;
    margin-bottom: 8px;
}

.section-description {
    color: #aeb9ce;
    font-size: 15px;
    line-height: 1.6;
    margin-bottom: 15px;
}

.status-badge {
    display: inline-block;
    background: #263653;
    color: #b9caff;
    border: 1px solid #526da8;
    padding: 5px 12px;
    border-radius: 20px;
    font-size: 12px;
    font-weight: 600;
    margin-bottom: 12px;
}

.progress-container {
    background: #30394d;
    border-radius: 20px;
    height: 8px;
    width: 100%;
    margin-top: 12px;
}

.progress-fill {
    background: linear-gradient(90deg, #7c9cff, #b58cff);
    height: 8px;
    border-radius: 20px;
}

.small-label {
    color: #8f9bb2;
    font-size: 12px;
    margin-top: 8px;
}

</style>
""", unsafe_allow_html=True)

st.markdown(
    """
    <style>

    .main-title {
        font-size: 36px;
        font-weight: 700;
        margin-bottom: 5px;
    }

    .subtitle {
        font-size: 17px;
        margin-bottom: 25px;
    }

    .metric-card {
        padding: 20px;
        border-radius: 12px;
        border: 1px solid #ddd;
        text-align: center;
        min-height: 120px;
    }

    .metric-title {
        font-size: 15px;
        font-weight: 600;
    }

    .metric-value {
        font-size: 25px;
        font-weight: 700;
        margin-top: 10px;
    }

    </style>
    """,
    unsafe_allow_html=True
    )




# ---------------------------------------------------------
# PAGE CONFIGURATION
# ---------------------------------------------------------
# ---------------------------------------------------------
# CUSTOM CSS
# ---------------------------------------------------------

st.markdown("""
<style>

.stApp {
    background-color: #0E1117;
    color: #F5F5F5;
}

.main .block-container {
    max-width: 1400px;
    padding-top: 10px;
    padding-left: 5%;
    padding-right: 5%;
}

/* Cards */
div[data-testid="stMetric"] {
    background-color: #161B22;
    border: 1px solid #30363D;
    border-radius: 12px;
    padding: 15px;
}

/* ---------------------------------------------------------
   BUTTONS
--------------------------------------------------------- */

.stButton > button {
    background-color: #151922 !important;
    color: #F5F5F5 !important;
    border: 1px solid #374151 !important;
    border-radius: 14px !important;
    min-height: 58px !important;
    font-size: 15px !important;
    font-weight: 600 !important;
    transition: all 0.2s ease !important;
}

.stButton > button:hover {
    background-color: #1D2330 !important;
    border-color: #8B5CF6 !important;
    color: #FFFFFF !important;
    transform: translateY(-2px);
    box-shadow: 0 6px 18px rgba(139, 92, 246, 0.15) !important;
}

.stButton > button:active {
    transform: translateY(0px);
}

/* Active Navigation Button */

.stButton > button[kind="primary"] {
    background-color: #211A35 !important;
    color: #FFFFFF !important;
    border: 1px solid #8B5CF6 !important;
    box-shadow:
        0 0 0 1px #8B5CF6 inset,
        0 6px 18px rgba(139, 92, 246, 0.18) !important;
}

/* Headings */
h1, h2, h3 {
    color: #F5F5F5;
}

/* Muted text */
.stCaption {
    color: #9CA3AF;
}

    /* Sidebar */
    section[data-testid="stSidebar"] {
        background-color: #151922;
        border-right: 1px solid #262b36;
    }

    .sidebar-title {
        font-size: 22px;
        font-weight: 700;
        margin-bottom: 5px;
    }

    .sidebar-subtitle {
        font-size: 12px;
        color: #8b93a7;
        margin-bottom: 25px;
    }

    /* Main heading */
    .main-title {
        font-size: 38px;
        font-weight: 750;
        margin-bottom: 5px;
    }

    .main-subtitle {
        color: #9aa3b5;
        font-size: 16px;
        margin-bottom: 30px;
    }

   
    /* Cards */

.metric-card {
    padding: 20px;
    border-radius: 12px;
    border: 1px solid #ddd;
}

.metric-card.risk-card {
    border-left: 4px solid #f4a261;
    background: rgba(244, 162, 97, 0.08);
}

.metric-card.blocker-card {
    border-left: 4px solid #e9c46a;
    background: rgba(233, 196, 106, 0.08);
}

.metric-card.health-card {
    border-left: 4px solid #8abf8a;
    background: rgba(138, 191, 138, 0.08);
}

.metric-card.document-card {
    border-left: 4px solid #7aa2f7;
    background: rgba(122, 162, 247, 0.08);
}

.metric-title {
    color: #8f98aa;
    font-size: 14px;
    margin-bottom: 10px;
}

    .metric-value {
        font-size: 30px;
        font-weight: 700;
    }

    /* Section */
    .section-title {
        font-size: 22px;
        font-weight: 650;
        margin-top: 30px;
        margin-bottom: 12px;
    }

    .section-description {
        color: #8f98aa;
        font-size: 14px;
        margin-bottom: 15px;
    }

    /* Upload box */
    .upload-container {
        background-color: #151922;
        border: 1px solid #272d39;
        border-radius: 12px;
        padding: 20px;
    }

    /* RAG pipeline */
    .pipeline-container {
        background-color: #151922;
        border: 1px solid #272d39;
        border-radius: 12px;
        padding: 25px;
        margin-top: 10px;
    }

    .pipeline-step {
        text-align: center;
        padding: 15px;
        background-color: #1b202b;
        border-radius: 10px;
        border: 1px solid #2a303d;
    }

    .pipeline-name {
        font-weight: 600;
        font-size: 14px;
    }

    .pipeline-status {
        color: #7d8799;
        font-size: 12px;
        margin-top: 5px;
    }

    /* Info box */
    .info-box {
        background-color: #151922;
        border: 1px solid #272d39;
        border-radius: 12px;
        padding: 20px;
        color: #aeb6c5;
        line-height: 1.6;
    }
    
    .status-badge {
    display: inline-block;
    padding: 5px 10px;
    border-radius: 20px;
    font-size: 12px;
    font-weight: 600;
}

.status-badge.completed {
    color: #8abf8a;
    background: rgba(138, 191, 138, 0.12);
    border: 1px solid rgba(138, 191, 138, 0.25);
}

.status-badge.waiting {
    color: #8f98aa;
    background: rgba(143, 152, 170, 0.08);
    border: 1px solid rgba(143, 152, 170, 0.20);
}

.home-hero {
    position: relative;
    display: flex;
    align-items: center;
    justify-content: space-between;
    gap: 20px;
    min-height: 365px;
    overflow: hidden;
    padding: 28px 34px;
    margin: 0 0 17px;
    border: 1px solid rgba(119, 146, 220, .24);
    border-radius: 24px;
    background:
        radial-gradient(ellipse at 80% 53%, rgba(31, 135, 255, .18), transparent 35%),
        radial-gradient(ellipse at 51% 5%, rgba(161, 76, 255, .10), transparent 43%),
        radial-gradient(ellipse at 4% 16%, rgba(122, 70, 228, .18), transparent 43%),
        linear-gradient(118deg, #090d19 0%, #10172a 53%, #0b1120 100%);
    box-shadow: 0 15px 42px rgba(3, 7, 18, .34), inset 0 1px rgba(255,255,255,.045);
}

.home-hero:after {
    content: "";
    position: absolute;
    inset: auto -7% -82% 18%;
    height: 265px;
    border: 1px solid rgba(34, 211, 238, .12);
    border-radius: 50%;
    transform: rotate(-7deg);
    box-shadow: 0 -10px 42px rgba(34, 211, 238, .055), 0 -25px 65px rgba(139, 92, 246, .065);
    pointer-events: none;
}

.home-hero:before { content: ""; position: absolute; inset: 0; pointer-events: none; opacity: .5; background-image: radial-gradient(circle at 58% 19%, rgba(164,190,255,.8) 0 1px, transparent 1.5px), radial-gradient(circle at 75% 22%, rgba(80,220,255,.8) 0 1px, transparent 1.5px), radial-gradient(circle at 92% 34%, rgba(255,147,214,.7) 0 1px, transparent 1.5px), radial-gradient(circle at 65% 75%, rgba(117,143,255,.8) 0 1px, transparent 1.5px), radial-gradient(circle at 48% 61%, rgba(80,220,255,.6) 0 1px, transparent 1.5px); }
.home-hero-copy { position: relative; z-index: 1; flex: 1 1 58%; min-width: 0; }
.home-eyebrow { color: #91a5ca; font-size: 11px; font-weight: 700; letter-spacing: .19em; }
.home-eyebrow:before { content: ""; display: inline-block; width: 7px; height: 7px; margin: 0 9px 1px 0; border-radius: 50%; background: #22d3ee; box-shadow: 0 0 12px #22d3ee; }
.home-title { margin: 15px 0 12px; color: #f7f8ff; font-size: clamp(34px, 4.1vw, 51px); line-height: 1.08; font-weight: 780; letter-spacing: -.04em; }
.home-title span { background: linear-gradient(90deg, #c6a8ff 0%, #f49bd8 44%, #77d9ff 82%); -webkit-background-clip: text; background-clip: text; color: transparent; }
.home-subtitle { max-width: 570px; margin: 0; color: #b5c0d4; font-size: 14px; line-height: 1.7; }
.home-badges { display: flex; flex-wrap: wrap; gap: 8px; margin-top: 20px; }
.home-badge { display: inline-flex; align-items: center; gap: 7px; padding: 8px 11px; color: #d4def1; font-size: 11px; font-weight: 600; border: 1px solid rgba(139, 157, 205, .22); border-radius: 999px; background: rgba(17, 26, 46, .72); box-shadow: inset 0 1px rgba(255,255,255,.04); white-space: nowrap; }
.home-badge:nth-child(1) { border-color: rgba(181,140,255,.34); box-shadow: 0 0 15px rgba(139,92,246,.11); }
.home-badge:nth-child(2) { border-color: rgba(255,111,145,.32); }
.home-badge:nth-child(3) { border-color: rgba(34,211,238,.30); }
.home-badge:nth-child(4) { border-color: rgba(126,214,165,.28); }
.home-badge-icon { font-size: 13px; line-height: 1; }
.home-visual { position: relative; z-index: 1; flex: 0 1 47%; min-width: 440px; height: 330px; }
.home-scene-lines { position: absolute; inset: 0; width: 100%; height: 100%; overflow: visible; }
.home-scene-lines path { stroke: rgba(82, 185, 255, .23); stroke-width: 1; fill: none; }
.home-scene-lines circle { fill: #59d9ff; filter: drop-shadow(0 0 5px #59d9ff); }
.home-orbit { position: absolute; left: 50%; bottom: 8px; width: 205px; height: 205px; border: 1px solid rgba(83, 176, 255, .25); border-radius: 50%; transform: translateX(-50%) scaleY(.34); box-shadow: 0 0 42px rgba(63, 146, 255, .22), inset 0 0 25px rgba(34, 211, 238, .12); }
.home-orbit:before, .home-orbit:after { content: ""; position: absolute; inset: 12px; border: 1px solid rgba(34, 211, 238, .25); border-radius: 50%; }
.home-orbit:after { inset: 26px; border-color: rgba(181, 140, 255, .27); }
.home-platform { position: absolute; z-index: 2; left: 50%; bottom: 16px; width: 168px; height: 20px; border: 1px solid rgba(63, 210, 255, .75); border-radius: 50%; transform: translateX(-50%); background: radial-gradient(ellipse, rgba(71, 197, 255, .38), rgba(62, 103, 255, .09) 57%, transparent 72%); box-shadow: 0 0 24px rgba(44, 175, 255, .42), 0 0 55px rgba(93, 88, 255, .2); }
.home-robot { position: absolute; z-index: 3; left: 50%; top: 91px; width: 170px; height: 222px; transform: translateX(-50%); filter: drop-shadow(0 10px 17px rgba(72, 148, 255, .23)); }
.robot-antenna { position: absolute; top: 4px; left: 83px; width: 4px; height: 18px; border-radius: 4px; background: linear-gradient(#91eeff,#6b8eff); box-shadow: 0 0 9px rgba(65,211,255,.7); }
.robot-antenna:before { content: ""; position: absolute; top: -5px; left: -4px; width: 12px; height: 12px; border-radius: 50%; background: #8af0ff; box-shadow: 0 0 12px #39cfff, 0 0 23px rgba(92,113,255,.75); }
.robot-head { position: absolute; top: 18px; left: 30px; width: 110px; height: 81px; border: 2px solid #b9e9ff; border-radius: 29px; background: linear-gradient(145deg,#ffffff 0%,#e4f2ff 58%,#a9d8f6 100%); box-shadow: inset 0 2px 5px rgba(255,255,255,.9), 0 0 21px rgba(70,172,255,.26); }
.robot-face { position: absolute; inset: 13px 11px 12px; border: 1px solid rgba(87,190,255,.55); border-radius: 18px; background: linear-gradient(145deg,#122039,#1b3152); box-shadow: inset 0 0 14px rgba(32,149,255,.18); }
.robot-eye { position: absolute; top: 14px; width: 13px; height: 9px; border-radius: 8px; background: #67e7ff; box-shadow: 0 0 7px #40d7ff, 0 0 14px rgba(49,160,255,.9); }
.robot-eye-left { left: 14px; }.robot-eye-right { right: 14px; }
.robot-mouth { position: absolute; left: 50%; bottom: 8px; width: 23px; height: 4px; border-radius: 4px; transform: translateX(-50%); background: #94cfff; box-shadow: 0 0 7px rgba(66,174,255,.75); }
.robot-ear { position: absolute; top: 27px; width: 9px; height: 26px; border: 2px solid #a7e4ff; border-radius: 7px; background: linear-gradient(#e8f8ff,#94ccef); }
.robot-ear-left { left: -7px; }.robot-ear-right { right: -7px; }
.robot-neck { position: absolute; top: 96px; left: 72px; width: 26px; height: 14px; border-radius: 5px; background: linear-gradient(90deg,#9ac9e9,#f5fcff,#89bfdf); }
.robot-body { position: absolute; top: 105px; left: 45px; width: 80px; height: 70px; border: 2px solid #b6e6ff; border-radius: 24px 24px 20px 20px; background: linear-gradient(145deg,#ffffff,#dceeff 63%,#9acbe9); box-shadow: inset 0 2px 6px rgba(255,255,255,.8), 0 0 14px rgba(75,174,255,.2); }
.robot-ai-badge { position: absolute; top: 17px; left: 20px; display: grid; place-items: center; width: 38px; height: 24px; border: 1px solid #79e6ff; border-radius: 8px; color: #d9fbff; background: linear-gradient(135deg,#176e9f,#303e9c); box-shadow: 0 0 10px rgba(52,202,255,.48); font-size: 11px; font-weight: 800; letter-spacing: .1em; }
.robot-body-light { position: absolute; bottom: 10px; left: 30px; width: 16px; height: 3px; border-radius: 4px; background: #59dfff; box-shadow: 0 0 7px #42d3ff; }
.robot-arm { position: absolute; z-index: 1; width: 14px; height: 47px; border: 2px solid #a8ddf8; border-radius: 10px; background: linear-gradient(90deg,#a5cee8,#f8fdff,#9cc8e6); }
.robot-arm:after { content: ""; position: absolute; bottom: -9px; left: -4px; width: 18px; height: 13px; border: 1px solid #83dcff; border-radius: 7px; background: linear-gradient(#f4fcff,#a5daf4); }
.robot-arm-left { top: 112px; left: 29px; transform: rotate(14deg); }
.robot-arm-right { top: 100px; right: 27px; height: 42px; transform: rotate(-29deg); }
.robot-leg { position: absolute; top: 168px; width: 16px; height: 29px; border-radius: 8px; background: linear-gradient(90deg,#a5cee8,#f8fdff,#9cc8e6); }
.robot-leg-left { left: 58px; }.robot-leg-right { right: 58px; }
.robot-foot { position: absolute; top: 191px; width: 26px; height: 12px; border: 1px solid #a9e3ff; border-radius: 9px 9px 6px 6px; background: linear-gradient(#faffff,#97cbea); box-shadow: 0 0 9px rgba(48,193,255,.3); }
.robot-foot-left { left: 48px; }.robot-foot-right { right: 48px; }
.home-float-card { position: absolute; z-index: 4; width: 144px; min-height: 68px; padding: 9px 10px; border: 1px solid rgba(122, 166, 231, .27); border-radius: 12px; color: #dce8fc; background: linear-gradient(145deg, rgba(21, 31, 53, .88), rgba(11, 18, 34, .82)); box-shadow: 0 8px 21px rgba(0,0,0,.19), inset 0 1px rgba(255,255,255,.055); backdrop-filter: blur(8px); }
.home-float-title { display: flex; align-items: center; gap: 6px; margin-bottom: 7px; color: #d9e4f7; font-size: 9px; font-weight: 700; letter-spacing: .02em; }
.home-float-title svg { width: 13px; height: 13px; flex: none; }
.home-insights { top: 18px; left: 2px; border-color: rgba(181,140,255,.34); }
.home-health { top: 11px; right: 0; border-color: rgba(34,211,238,.31); }
.home-risk { bottom: 38px; left: 0; border-color: rgba(255,111,145,.31); }
.home-recommendations { right: 0; bottom: 32px; border-color: rgba(255,159,67,.30); }
.home-mini-bars { display: flex; align-items: end; gap: 4px; height: 18px; padding-left: 2px; }
.home-mini-bars i { display: block; width: 8px; border-radius: 3px 3px 1px 1px; background: linear-gradient(#5de1ff,#6578ff); box-shadow: 0 0 7px rgba(67,185,255,.25); }
.home-mini-bars i:nth-child(1) { height: 8px; }.home-mini-bars i:nth-child(2) { height: 13px; }.home-mini-bars i:nth-child(3) { height: 10px; }.home-mini-bars i:nth-child(4) { height: 17px; }.home-mini-bars i:nth-child(5) { height: 14px; }
.home-health-row { display: flex; align-items: center; gap: 8px; }
.home-donut { width: 27px; height: 27px; flex: none; border: 4px solid #36d5f5; border-right-color: #9778ff; border-bottom-color: rgba(106,132,185,.27); border-radius: 50%; box-shadow: 0 0 9px rgba(34,211,238,.25); }
.home-indicators { flex: 1; display: grid; gap: 5px; }
.home-indicators i { display: block; height: 3px; border-radius: 8px; background: linear-gradient(90deg, #56d9f4 0 68%, rgba(129,151,191,.22) 68%); }
.home-indicators i:nth-child(2) { background: linear-gradient(90deg, #a181ff 0 82%, rgba(129,151,191,.22) 82%); }
.home-indicators i:nth-child(3) { background: linear-gradient(90deg, #55c9ff 0 55%, rgba(129,151,191,.22) 55%); }
.home-risk-row, .home-recommendation-row { display: flex; align-items: center; gap: 8px; }
.home-risk-mark, .home-bulb-mark { display: grid; place-items: center; width: 24px; height: 24px; flex: none; border-radius: 8px; background: rgba(255,111,145,.12); }
.home-risk-mark svg, .home-bulb-mark svg { width: 15px; height: 15px; }
.home-bulb-mark { background: rgba(255,159,67,.12); }
.home-indicators-risk i { background: linear-gradient(90deg, #ff789e 0 72%, rgba(129,151,191,.22) 72%); }
.home-indicators-risk i:nth-child(2) { background: linear-gradient(90deg, #ffac73 0 48%, rgba(129,151,191,.22) 48%); }
.home-indicators-risk i:nth-child(3) { background: linear-gradient(90deg, #a181ff 0 65%, rgba(129,151,191,.22) 65%); }
.home-indicators-rec i { background: linear-gradient(90deg, #ffbd75 0 76%, rgba(129,151,191,.22) 76%); }
.home-indicators-rec i:nth-child(2) { background: linear-gradient(90deg, #61d8f4 0 58%, rgba(129,151,191,.22) 58%); }
.home-indicators-rec i:nth-child(3) { background: linear-gradient(90deg, #a181ff 0 84%, rgba(129,151,191,.22) 84%); }

[class*="st-key-nav_row"] button { min-height: 56px !important; border: 1px solid rgba(116, 144, 203, .28) !important; border-radius: 14px !important; background: linear-gradient(145deg, rgba(20, 29, 49, .98), rgba(12, 19, 34, .98)) !important; color: #cbd7ed !important; box-shadow: inset 0 1px rgba(255,255,255,.045), 0 5px 15px rgba(0,0,0,.14) !important; }
[class*="st-key-nav_row"] button:hover { border-color: rgba(74, 202, 255, .65) !important; background: linear-gradient(145deg, #1b2741, #121e35) !important; box-shadow: 0 0 18px rgba(66, 151, 255, .14) !important; }
[class*="st-key-nav_row"] button[kind="primary"] { border-color: rgba(168, 122, 255, .85) !important; color: #ffffff !important; background: linear-gradient(125deg, rgba(105, 73, 185, .52), rgba(28, 111, 164, .42)) !important; box-shadow: 0 0 0 1px rgba(164, 126, 255, .25) inset, 0 0 19px rgba(126, 89, 255, .23) !important; }
[class*="st-key-nav_row"] button:before { display: inline-block; margin-right: 7px; color: #67dcff; filter: drop-shadow(0 0 4px rgba(82,204,255,.48)); }
.st-key-nav_row1_0 button:before { content: "🏠"; }.st-key-nav_row1_1 button:before { content: "📁"; }.st-key-nav_row1_2 button:before { content: "🔗"; }.st-key-nav_row1_3 button:before { content: "🎯"; }.st-key-nav_row1_4 button:before { content: "⚠️"; }
.st-key-nav_row2_0 button:before { content: "☷"; }.st-key-nav_row2_1 button:before { content: "♥"; color: #ff83b1; }.st-key-nav_row2_2 button:before { content: "▤"; }.st-key-nav_row2_3 button:before { content: "▥"; }.st-key-nav_row2_4 button:before { content: "🤖"; }

.home-card { position: relative; height: 100%; min-height: 186px; padding: 21px 23px; overflow: hidden; border: 1px solid rgba(123, 143, 192, .25); border-radius: 19px; background: linear-gradient(145deg, rgba(24, 31, 51, .92), rgba(14, 20, 35, .88)); box-shadow: 0 10px 25px rgba(3, 7, 18, .19), inset 0 1px rgba(255,255,255,.045); text-align: center; }
.home-card:before { content: ""; position: absolute; top: 0; left: 22%; width: 56%; height: 2px; background: linear-gradient(90deg, transparent, var(--home-accent), transparent); box-shadow: 0 0 13px var(--home-accent); }
.home-card:hover { border-color: color-mix(in srgb, var(--home-accent) 58%, transparent); box-shadow: 0 13px 31px rgba(25, 35, 78, .25), 0 0 22px color-mix(in srgb, var(--home-accent) 12%, transparent); }
.home-card-doc { --home-accent: #a98aff; border-color: rgba(169,138,255,.30); }
.home-card-ai { --home-accent: #55d9ed; border-color: rgba(85,217,237,.28); }
.home-card-risk { --home-accent: #ff7fa8; border-color: rgba(255,127,168,.28); }
.home-card-icon { display: grid; place-items: center; width: 56px; height: 56px; margin: 0 auto 11px; border: 1px solid color-mix(in srgb, var(--home-accent) 52%, transparent); border-radius: 50%; background: radial-gradient(circle, color-mix(in srgb, var(--home-accent) 20%, transparent), rgba(10,16,30,.78)); box-shadow: 0 0 20px color-mix(in srgb, var(--home-accent) 20%, transparent), inset 0 0 13px color-mix(in srgb, var(--home-accent) 9%, transparent); }
.home-card-icon svg { width: 25px; height: 25px; }
.home-icon-document { position: relative; width: 20px; height: 25px; border: 2px solid #e9ddff; border-radius: 4px; background: linear-gradient(145deg,#fff,#d9c7ff); box-shadow: 0 0 10px rgba(183,143,255,.55); }
.home-icon-document:before { content: ""; position: absolute; top: -2px; right: -2px; width: 8px; height: 8px; border-left: 2px solid #b294ff; border-bottom: 2px solid #b294ff; border-radius: 0 2px 0 3px; background: #d5c3ff; }
.home-icon-document i { position: relative; display: block; width: 9px; height: 2px; margin: 12px 0 0 4px; border-radius: 2px; background: #8866d5; }
.home-icon-document i + i { width: 11px; margin-top: 4px; }
.home-icon-neural { position: relative; width: 29px; height: 29px; background: linear-gradient(30deg,transparent 48%,rgba(99,232,255,.95) 49% 52%,transparent 53%), linear-gradient(150deg,transparent 47%,rgba(144,145,255,.95) 48% 51%,transparent 52%), linear-gradient(90deg,transparent 48%,rgba(99,232,255,.8) 49% 51%,transparent 52%); filter: drop-shadow(0 0 5px rgba(70,205,255,.45)); }
.home-icon-neural i { position: absolute; z-index: 1; width: 7px; height: 7px; border: 1px solid #d7fcff; border-radius: 50%; background: #42d8ff; box-shadow: 0 0 7px #35c9ff; }
.home-icon-neural i:nth-child(1) { top: 1px; left: 11px; }.home-icon-neural i:nth-child(2) { top: 10px; left: 1px; }.home-icon-neural i:nth-child(3) { top: 10px; right: 1px; }.home-icon-neural i:nth-child(4) { bottom: 1px; left: 4px; }.home-icon-neural i:nth-child(5) { bottom: 1px; right: 4px; }
.home-icon-shield { position: relative; display: grid; place-items: center; width: 25px; height: 29px; clip-path: polygon(50% 0,95% 17%,89% 65%,50% 100%,11% 65%,5% 17%); color: #fff2f6; background: linear-gradient(145deg,#ff9bbb,#e55b91 58%,#aa477e); filter: drop-shadow(0 0 7px rgba(255,111,145,.6)); font-size: 15px; font-weight: 800; }
.home-card-title { margin-bottom: 7px; color: #f2f5ff; font-size: 16px; font-weight: 680; }
.home-card-text { max-width: 290px; margin: auto; color: #9eabc3; font-size: 12px; line-height: 1.55; }
.home-footer { display: flex; flex-wrap: wrap; justify-content: center; align-items: center; gap: 0; margin-top: 20px; padding: 15px 10px 4px; text-align: center; color: #93a1b9; font-size: 11px; letter-spacing: .025em; border-top: 1px solid rgba(132, 149, 190, .16); }
.home-footer-item { display: inline-flex; align-items: center; gap: 6px; padding: 0 13px; }
.home-footer-item:not(:last-child) { border-right: 1px solid rgba(117,139,183,.25); }
.home-footer-item svg { width: 13px; height: 13px; color: #65d7f3; }

@media (max-width: 850px) {
    .home-hero { min-height: 0; padding: 28px 24px; }
    .home-visual { flex-basis: 44%; min-width: 375px; transform: scale(.88); transform-origin: center right; }
    .home-title { font-size: 37px; }
    .home-badge { font-size: 10px; padding: 7px 9px; }
}

@media (max-width: 640px) {
    .home-hero { display: block; padding: 23px 18px 12px; }
    .home-visual { min-width: 0; width: 100%; height: 280px; margin-top: 8px; transform: scale(.82); transform-origin: top center; }
    .home-title { font-size: 34px; }
    .home-badges { gap: 6px; margin-top: 15px; }
    .home-badge { white-space: normal; }
    .home-footer-item { padding: 6px 9px; }
    .home-footer-item:not(:last-child) { border-right: 0; }
}
</style>
""", unsafe_allow_html=True)



# ---------------------------------------------------------
# SIDEBAR
# ---------------------------------------------------------
import streamlit as st

# ---------------- SIDEBAR DESIGN ----------------
st.markdown(
    """
    <style>
    /* Sidebar width */
    section[data-testid="stSidebar"] {
        width: 290px !important;
        background: linear-gradient(180deg, #111827 0%, #0f172a 100%);
        border-right: 1px solid #263244;
    }

    /* Sidebar padding */
    section[data-testid="stSidebar"] > div {
        padding: 28px 18px;
    }

    /* Sidebar title */
    .sidebar-title {
        font-size: 26px;
        font-weight: 800;
        color: #ffffff;
        margin-bottom: 4px;
    }

    .sidebar-subtitle {
        font-size: 14px;
        color: #94a3b8;
        margin-bottom: 28px;
    }

    .sidebar-line {
        height: 1px;
        background: #334155;
        margin: 18px 0 24px 0;
    }

    /* Radio button labels */
    section[data-testid="stSidebar"] label {
        color: #e5e7eb !important;
        font-size: 15px !important;
        font-weight: 500 !important;
        padding: 9px 12px !important;
        border-radius: 10px !important;
        margin-bottom: 5px !important;
        transition: all 0.2s ease-in-out;
    }

    section[data-testid="stSidebar"] label:hover {
        background: #1e293b !important;
        color: #ffffff !important;
    }

    /* Hide radio circles */
    section[data-testid="stSidebar"] div[role="radiogroup"] > label > div:first-child {
        display: none !important;
    }

    /* Radio group spacing */
    section[data-testid="stSidebar"] div[role="radiogroup"] {
        gap: 7px !important;
    }

    /* Selected navigation item */
    section[data-testid="stSidebar"] div[role="radiogroup"] label:has(input:checked) {
        background: linear-gradient(90deg, #2563eb, #1d4ed8) !important;
        color: white !important;
        box-shadow: 0 5px 14px rgba(37, 99, 235, 0.25);
    }

    </style>
    """,
    unsafe_allow_html=True
)



# ---------------------------------------------------------
# HOME PAGE
# ---------------------------------------------------------

def show_home():

    st.html("""
<style>
.stApp {
    background:
        radial-gradient(ellipse at 84% 22%, rgba(48, 113, 194, .12), transparent 31%),
        radial-gradient(ellipse at 8% 4%, rgba(106, 65, 176, .13), transparent 28%),
        #080b14 !important;
}
.main .block-container { max-width: 1460px; padding-top: 20px; padding-bottom: 20px; }
.home-hero, .home-hero * { box-sizing: border-box; }
</style>
<section class="home-hero">
    <div class="home-hero-copy">
        <div class="home-eyebrow">AI PROJECT INTELLIGENCE PLATFORM</div>
        <h1 class="home-title">AI Project Intelligence<br><span>&amp; Risk Advisor</span></h1>
        <p class="home-subtitle">AI-powered project monitoring, risk analysis, and decision support</p>
        <div class="home-badges">
            <span class="home-badge"><span class="home-badge-icon">✨</span>Smarter Insights</span>
            <span class="home-badge"><span class="home-badge-icon">⚡</span>Identify Risks Early</span>
            <span class="home-badge"><span class="home-badge-icon">📊</span>Track Progress</span>
            <span class="home-badge"><span class="home-badge-icon">🎯</span>Take Informed Decisions</span>
        </div>
    </div>
    <div class="home-visual" aria-hidden="true">
        <svg class="home-scene-lines" viewBox="0 0 480 330" preserveAspectRatio="none">
            <path d="M120 75 C175 74 177 116 215 126"/><path d="M365 70 C315 72 305 110 265 126"/>
            <path d="M115 255 C164 248 178 218 214 207"/><path d="M365 251 C319 246 300 218 267 207"/>
            <circle cx="120" cy="75" r="2"/><circle cx="365" cy="70" r="2"/><circle cx="115" cy="255" r="2"/><circle cx="365" cy="251" r="2"/>
        </svg>
        <div class="home-orbit"></div>
        <div class="home-float-card home-insights">
            <div class="home-float-title"><svg viewBox="0 0 16 16" fill="none"><path d="M2 13V3m0 10h12M5 10V7m3 3V4m3 6V6" stroke="#b58cff" stroke-width="1.5" stroke-linecap="round"/></svg>Project Insights</div>
            <div class="home-mini-bars"><i></i><i></i><i></i><i></i><i></i></div>
        </div>
        <div class="home-float-card home-health">
            <div class="home-float-title"><svg viewBox="0 0 16 16" fill="none"><path d="M8 1.8 13 3.7v3.5c0 3-2 5.3-5 7-3-1.7-5-4-5-7V3.7l5-1.9Z" stroke="#55d9ed" stroke-width="1.3"/><path d="m5.5 7.8 1.6 1.6 3.3-3.5" stroke="#55d9ed" stroke-width="1.3" stroke-linecap="round" stroke-linejoin="round"/></svg>Project Health</div>
            <div class="home-health-row"><div class="home-donut"></div><div class="home-indicators"><i></i><i></i><i></i></div></div>
        </div>
        <div class="home-float-card home-risk">
            <div class="home-float-title"><svg viewBox="0 0 16 16" fill="none"><path d="M8 1.5 14 4v4c0 3.3-2.5 5.4-6 6.7C4.5 13.4 2 11.3 2 8V4l6-2.5Z" stroke="#ff789e" stroke-width="1.3"/><path d="M8 5v3m0 2.1h.01" stroke="#ff789e" stroke-width="1.3" stroke-linecap="round"/></svg>Risk Analysis</div>
            <div class="home-risk-row"><div class="home-risk-mark"><svg viewBox="0 0 16 16" fill="none"><path d="M8 1.5 14 4v4c0 3.3-2.5 5.4-6 6.7C4.5 13.4 2 11.3 2 8V4l6-2.5Z" stroke="#ff789e" stroke-width="1.3"/></svg></div><div class="home-indicators home-indicators-risk"><i></i><i></i><i></i></div></div>
        </div>
        <div class="home-float-card home-recommendations">
            <div class="home-float-title"><svg viewBox="0 0 16 16" fill="none"><path d="M5 10.5c-.2-1.2-1.7-2-1.7-4A4.7 4.7 0 0 1 8 1.8a4.7 4.7 0 0 1 4.7 4.7c0 2-1.5 2.8-1.7 4H5Z" stroke="#ffb65c" stroke-width="1.25"/><path d="M6 12.5h4m-3.3 2h2.6" stroke="#ffb65c" stroke-width="1.25" stroke-linecap="round"/></svg>Recommendations</div>
            <div class="home-recommendation-row"><div class="home-bulb-mark"><svg viewBox="0 0 16 16" fill="none"><path d="M5 10.5c-.2-1.2-1.7-2-1.7-4A4.7 4.7 0 0 1 8 1.8a4.7 4.7 0 0 1 4.7 4.7c0 2-1.5 2.8-1.7 4H5Z" stroke="#ffb65c" stroke-width="1.25"/><path d="M6 12.5h4m-3.3 2h2.6" stroke="#ffb65c" stroke-width="1.25" stroke-linecap="round"/></svg></div><div class="home-indicators home-indicators-rec"><i></i><i></i><i></i></div></div>
        </div>
        <div class="home-robot">
            <div class="robot-antenna"></div>
            <div class="robot-ear robot-ear-left"></div>
            <div class="robot-ear robot-ear-right"></div>
            <div class="robot-head">
                <div class="robot-face">
                    <span class="robot-eye robot-eye-left"></span>
                    <span class="robot-eye robot-eye-right"></span>
                    <span class="robot-mouth"></span>
                </div>
            </div>
            <div class="robot-arm robot-arm-left"></div>
            <div class="robot-arm robot-arm-right"></div>
            <div class="robot-neck"></div>
            <div class="robot-body">
                <div class="robot-ai-badge">AI</div>
                <div class="robot-body-light"></div>
            </div>
            <div class="robot-leg robot-leg-left"></div>
            <div class="robot-leg robot-leg-right"></div>
            <div class="robot-foot robot-foot-left"></div>
            <div class="robot-foot robot-foot-right"></div>
        </div>
        <div class="home-platform"></div>
    </div>
</section>
""")


def _build_rag_context(results):
    context = "\n\n".join(
        result.get("text", "")
        for result in results
    )

    analysis_context = [
        ("PROJECT SCOPE ANALYSIS", st.session_state.get("scope_analysis", "")),
        ("RISK ANALYSIS", st.session_state.get("risk_analysis", "")),
        ("BLOCKER & ACTION ANALYSIS", st.session_state.get("blocker_analysis", "")),
        ("PROJECT HEALTH ANALYSIS", st.session_state.get("health_analysis", "")),
    ]

    for heading, analysis in analysis_context:
        if analysis:
            context += f"\n\n===== {heading} =====\n" + str(analysis)

    return context


def _prepare_project_scope_entry():
    for analysis_key in ("scope_analysis", "risk_analysis", "blocker_analysis"):
        if not st.session_state.get(analysis_key):
            st.session_state[f"{analysis_key}_attempted"] = False


# ---------------------------------------------------------
# SESSION STATE
# ---------------------------------------------------------

if "selected_page" not in st.session_state:
    st.session_state.selected_page = "Home"

if "processed_documents" not in st.session_state:
    st.session_state.processed_documents = []

if "scope_analysis" not in st.session_state:
    st.session_state.scope_analysis = ""
if "scope_analysis_attempted" not in st.session_state:
    st.session_state.scope_analysis_attempted = bool(st.session_state.get("scope_analysis"))

if "risk_analysis" not in st.session_state:
    st.session_state.risk_analysis = ""
if "risk_analysis_attempted" not in st.session_state:
    st.session_state.risk_analysis_attempted = bool(st.session_state.get("risk_analysis"))

if "blocker_analysis" not in st.session_state:
    st.session_state.blocker_analysis = ""
if "blocker_analysis_attempted" not in st.session_state:
    st.session_state.blocker_analysis_attempted = bool(st.session_state.get("blocker_analysis"))

if "health_analysis" not in st.session_state:
    st.session_state.health_analysis = ""

if "documentation" not in st.session_state:
    st.session_state.documentation = ""

selected_page = st.session_state.selected_page

if selected_page == "Project Health":

    st.html("""
    <div class="module-card" style="--accent:#7ed6a5;--progress:78%;">
        <div class="module-icon">💚</div>
        <div class="module-title">Project Health</div>
        <span class="module-badge">Health Intelligence</span>
        <div class="module-description">
            Evaluate project progress, delivery confidence, risks,
            blockers, and overall health using the processed project documents.
        </div>
        <div class="module-line"><span></span></div>
        <div class="module-caption">
            Scope • Progress • Timeline • Risks • Blockers • Delivery Confidence
        </div>
    </div>
    """)

    st.markdown("### 💚 Analyze Project Health")

    if st.button(
        "💚 Analyze Project Health",
        type="primary",
        use_container_width=True,
        key="analyze_project_health_button"
    ):

        if not st.session_state.get("processed_documents"):

            st.warning(
                "No project documents have been processed yet. "
                "Please go to Documents, upload your project files, "
                "and process them first."
            )

        else:

            health_chunks = retrieve_documents(
                "project health overall score scope clarity progress "
                "timeline delivery risks blockers delivery confidence",
                top_k=10
            )

            if health_chunks:

                with st.spinner("Analyzing project health..."):
                    st.session_state.health_analysis = analyze_project_health(
                        health_chunks
                    )

                st.success("Project health analysis completed.")

            else:

                st.warning(
                    "No relevant project information found "
                    "in the uploaded project documents."
                )

    if st.session_state.get("health_analysis"):

        health_text = str(st.session_state.health_analysis)

        overall_score_match = re.search(
            r"(?im)^\s*Score:\s*(.+?)\s*$",
            health_text
        )
        overall_status_match = re.search(
            r"(?im)^\s*Status:\s*(.+?)\s*$",
            health_text
        )
        overall_score = (
            overall_score_match.group(1).strip()
            if overall_score_match
            else "Not available"
        )
        overall_status = (
            overall_status_match.group(1).strip()
            if overall_status_match
            else "Not available"
        )
        overall_numeric_match = re.search(
            r"^\s*Score:\s*(\d{1,3})\s*/\s*100\b",
            health_text,
            re.IGNORECASE | re.MULTILINE
        )
        overall_progress = (
            min(100, max(0, int(overall_numeric_match.group(1))))
            if overall_numeric_match
            else 0
        )

        st.markdown("### 📋 Project Health Analysis")
        st.caption(
            "AI-powered assessment of project health, progress, risks and delivery confidence."
        )
        st.markdown(
            f"""
            <div style="
                display:flex;
                justify-content:space-between;
                align-items:center;
                gap:24px;
                padding:24px 28px;
                margin:8px 0 24px 0;
                border:1px solid #34415b;
                border-radius:18px;
                background:radial-gradient(circle at 5% 0%, rgba(124,156,255,.16), transparent 42%),
                           linear-gradient(135deg, #171d2b 0%, #20283a 100%);
                box-shadow:0 8px 26px rgba(50,80,150,.14);
            ">
                <div style="flex:1; min-width:0;">
                    <div style="color:#aeb9ce; font-size:14px; font-weight:600;">
                        Overall Project Health
                    </div>
                    <div style="color:#ffffff; font-size:36px; font-weight:800; margin:5px 0 14px 0;">
                        {html.escape(overall_score)}
                    </div>
                    <div style="height:8px; background:#30394d; border-radius:99px; overflow:hidden;">
                        <div style="height:100%; width:{overall_progress}%; background:linear-gradient(90deg,#7c9cff,#b58cff); border-radius:99px;"></div>
                    </div>
                </div>
                <div style="display:flex; align-items:center; gap:14px;">
                    <div style="
                        min-width:150px;
                        padding:14px 18px;
                        border:1px solid rgba(255,255,255,.10);
                        border-radius:14px;
                        background:rgba(10,14,22,.30);
                    ">
                        <div style="color:#8f9bb2; font-size:12px; text-transform:uppercase; letter-spacing:.08em;">
                            Status
                        </div>
                        <div style="color:#f5f5f5; font-size:18px; font-weight:700; margin-top:6px;">
                            {html.escape(overall_status)}
                        </div>
                    </div>
                    <div style="
                        width:48px;
                        height:48px;
                        flex:0 0 48px;
                        display:flex;
                        align-items:center;
                        justify-content:center;
                        border-radius:14px;
                        background:rgba(124,156,255,.14);
                        box-shadow:0 0 20px rgba(124,156,255,.18);
                    ">
                        <svg width="25" height="25" viewBox="0 0 24 24" fill="none" aria-hidden="true">
                            <path d="M4 19V5M4 19H21" stroke="#7c9cff" stroke-width="1.8" stroke-linecap="round"/>
                            <path d="m7 15 4-4 3 2 5-6" stroke="#7c9cff" stroke-width="1.8" stroke-linecap="round" stroke-linejoin="round"/>
                            <path d="M16 7h3v3" stroke="#7c9cff" stroke-width="1.8" stroke-linecap="round" stroke-linejoin="round"/>
                        </svg>
                    </div>
                </div>
            </div>
            """,
            unsafe_allow_html=True
        )

        st.markdown("### 📊 Dimension Scores")
        st.caption("Detailed scores across key project health dimensions")

        dimensions = [
            (
                "Scope Clarity",
                '<svg width="24" height="24" viewBox="0 0 24 24" fill="none" aria-hidden="true"><circle cx="12" cy="12" r="8" stroke="#7c9cff" stroke-width="1.8"/><circle cx="12" cy="12" r="4" stroke="#7c9cff" stroke-width="1.8"/><circle cx="12" cy="12" r="1" fill="#7c9cff"/><path d="M15 9 21 3m-4 0h4v4" stroke="#7c9cff" stroke-width="1.8" stroke-linecap="round" stroke-linejoin="round"/></svg>',
                "#7c9cff",
                "#202b47",
                "rgba(124,156,255,.18)"
            ),
            (
                "Progress",
                '<svg width="24" height="24" viewBox="0 0 24 24" fill="none" aria-hidden="true"><path d="M4 20V4m0 16h17" stroke="#b58cff" stroke-width="1.8" stroke-linecap="round"/><rect x="7" y="12" width="3" height="5" rx=".7" fill="#b58cff"/><rect x="12" y="9" width="3" height="8" rx=".7" fill="#b58cff"/><rect x="17" y="5" width="3" height="12" rx=".7" fill="#b58cff"/></svg>',
                "#b58cff",
                "#302543",
                "rgba(181,140,255,.18)"
            ),
            (
                "Timeline / Delivery",
                '<svg width="24" height="24" viewBox="0 0 24 24" fill="none" aria-hidden="true"><rect x="4" y="6" width="16" height="14" rx="2" stroke="#7ed6a5" stroke-width="1.8"/><path d="M8 4v4m8-4v4M4 10h16" stroke="#7ed6a5" stroke-width="1.8" stroke-linecap="round"/><path d="m9 15 2 2 4-4" stroke="#7ed6a5" stroke-width="1.8" stroke-linecap="round" stroke-linejoin="round"/></svg>',
                "#7ed6a5",
                "#20392f",
                "rgba(126,214,165,.16)"
            ),
            (
                "Risk Management",
                '<svg width="24" height="24" viewBox="0 0 24 24" fill="none" aria-hidden="true"><path d="M12 3.8 21 20H3L12 3.8Z" stroke="#ff9f43" stroke-width="1.8" stroke-linejoin="round"/><path d="M12 9v5m0 3h.01" stroke="#ff9f43" stroke-width="1.8" stroke-linecap="round"/></svg>',
                "#ff9f43",
                "#3e2d1d",
                "rgba(255,159,67,.17)"
            ),
            (
                "Blocker Management",
                '<svg width="24" height="24" viewBox="0 0 24 24" fill="none" aria-hidden="true"><circle cx="9" cy="8" r="3" stroke="#ff6f91" stroke-width="1.8"/><path d="M3.5 19c.4-3.1 2.2-5 5.5-5s5.1 1.9 5.5 5" stroke="#ff6f91" stroke-width="1.8" stroke-linecap="round"/><path d="M16 5.5a3 3 0 0 1 0 5.8m1 2.1c2.1.6 3.3 2.1 3.6 4.6" stroke="#ff6f91" stroke-width="1.8" stroke-linecap="round"/></svg>',
                "#ff6f91",
                "#402430",
                "rgba(255,111,145,.17)"
            ),
            (
                "Delivery Confidence",
                '<svg width="24" height="24" viewBox="0 0 24 24" fill="none" aria-hidden="true"><path d="M12 3 20 6v5.5c0 4.5-3.1 7.7-8 9.5-4.9-1.8-8-5-8-9.5V6l8-3Z" stroke="#38bdf8" stroke-width="1.8" stroke-linejoin="round"/><path d="m8.5 12 2.3 2.3 4.8-5" stroke="#38bdf8" stroke-width="1.8" stroke-linecap="round" stroke-linejoin="round"/></svg>',
                "#38bdf8",
                "#193545",
                "rgba(56,189,248,.17)"
            ),
        ]
        dimension_columns = st.columns(3)

        for index, (dimension_name, dimension_icon, accent, icon_background, icon_glow) in enumerate(dimensions):
            dimension_match = re.search(
                rf"(?im)^\s*{re.escape(dimension_name)}:\s*(.+?)\s*$",
                health_text
            )
            dimension_score = (
                dimension_match.group(1).strip()
                if dimension_match
                else "Not available"
            )
            dimension_numeric_match = re.search(r"\b(\d{1,3})\s*/\s*100\b", dimension_score)
            dimension_progress = (
                min(100, max(0, int(dimension_numeric_match.group(1))))
                if dimension_numeric_match
                else 0
            )
            with dimension_columns[index % 3]:
                st.markdown(
                    f"""
                    <div style="
                        min-height:132px;
                        padding:17px 18px;
                        margin:6px 0 12px 0;
                        border:1px solid #30394d;
                        border-left:3px solid {accent};
                        border-radius:15px;
                        background:linear-gradient(145deg,#171d2b 0%,#1d2535 100%);
                        box-shadow:0 6px 18px rgba(0,0,0,.15);
                    ">
                        <div style="display:flex; align-items:center; gap:13px; margin-bottom:10px;">
                            <div style="
                                width:48px;
                                height:48px;
                                flex:0 0 48px;
                                display:flex;
                                align-items:center;
                                justify-content:center;
                                border-radius:14px;
                                background:{icon_background};
                                box-shadow:0 0 18px {icon_glow};
                            ">
                                {dimension_icon}
                            </div>
                            <div style="font-size:13px; color:#aeb9ce; font-weight:600;">
                                {html.escape(dimension_name)}
                            </div>
                        </div>
                        <div style="font-size:23px; color:#ffffff; font-weight:750; margin:10px 0 13px 0;">
                            {html.escape(dimension_score)}
                        </div>
                        <div style="height:6px; background:#30394d; border-radius:99px; overflow:hidden;">
                            <div style="height:100%; width:{dimension_progress}%; background:{accent}; border-radius:99px;"></div>
                        </div>
                    </div>
                    """,
                    unsafe_allow_html=True
                )

        with st.expander("View Complete Project Health Analysis", expanded=True):
            compact_health_text = re.sub(
                r"(?m)^#{1,6}\s+(.+?)\s*$",
                r"**\1**",
                health_text
            )
            st.markdown(compact_health_text)

    else:

        st.info(
            "No project health analysis is available yet. "
            "Click Analyze Project Health."
        )

    st.divider()

    health_navigation_columns = st.columns([1, 2, 1])
    with health_navigation_columns[0]:
        if st.button("← Back to Home", key="health_back_home"):
            st.session_state.selected_page = "Home"
            st.rerun()
    with health_navigation_columns[2]:
        if st.button("Go to Documentation →", key="health_next_documentation", use_container_width=True):
            st.session_state.selected_page = "Documentation"
            st.rerun()

    st.stop()

# Fall back to Home when a stored navigation value has no page renderer.
if st.session_state.selected_page not in {
    "Home",
    "Dashboard",
    "RAG Pipeline",
    "Documents",
    "Project Scope",
    "Risks & Forecast",
    "Blockers & Actions",
    "Documentation",
    "AI Assistant",
}:
    st.session_state.selected_page = "Home"

# ---------------------------------------------------------
# SHOW HOME HEADER
# ---------------------------------------------------------


if st.session_state.selected_page == "Home":
    show_home()


# ---------------------------------------------------------
# TOP NAVIGATION
# ---------------------------------------------------------

if st.session_state.selected_page == "Home":

    menu_items_row1 = [
        ("Home", "Home"),
        ("Documents", "Documents"),
        ("RAG Pipeline", "RAG Pipeline"),
        ("Project Scope", "Project Scope"),
        ("Risks & Forecast", "Risks & Forecast"),
    ]

    menu_items_row2 = [
        ("Blockers & Actions", "Blockers & Actions"),
        ("Project Health", "Project Health"),
        ("Documentation", "Documentation"),
        ("Dashboard", "Dashboard"),
        ("AI Assistant", "AI Assistant"),
    ]

    # ROW 1

    nav_cols = st.columns(5)

    for i, (page_name, display_name) in enumerate(menu_items_row1):

        with nav_cols[i]:

            if st.button(
                display_name,
                key=f"nav_row1_{i}",
                use_container_width=True,
                type=(
                    "primary"
                    if st.session_state.selected_page == page_name
                    else "secondary"
                )
            ):
                if page_name == "Project Scope":
                    _prepare_project_scope_entry()
                st.session_state.selected_page = page_name
                st.rerun()

    # ROW 2

    nav_cols = st.columns(5)

    for i, (page_name, display_name) in enumerate(menu_items_row2):

        with nav_cols[i]:

            if st.button(
                display_name,
                key=f"nav_row2_{i}",
                use_container_width=True,
                type=(
                    "primary"
                    if st.session_state.selected_page == page_name
                    else "secondary"
                )
            ):
                st.session_state.selected_page = page_name
                st.rerun()


selected_page = st.session_state.selected_page


# ---------------------------------------------------------
# HOME FEATURE CARDS
# ---------------------------------------------------------

if selected_page == "Home":

    col1, col2, col3 = st.columns(3)

    with col1:
        st.html("""
<div class="home-card home-card-doc">
    <div class="home-card-icon"><div class="home-icon-document"><i></i><i></i></div></div>
    <div class="home-card-title">Document Intelligence</div>
    <div class="home-card-text">Upload project documents and extract meaningful project information.</div>
</div>
""")

    with col2:
        st.html("""
<div class="home-card home-card-ai">
    <div class="home-card-icon"><div class="home-icon-neural"><i></i><i></i><i></i><i></i><i></i></div></div>
    <div class="home-card-title">AI Analysis</div>
    <div class="home-card-text">Analyze project scope, progress, risks, blockers and delivery status.</div>
</div>
""")

    with col3:
        st.html("""
<div class="home-card home-card-risk">
    <div class="home-card-icon"><div class="home-icon-shield"><span>✓</span></div></div>
    <div class="home-card-title">Risk Intelligence</div>
    <div class="home-card-text">Identify project risks, blockers and actionable recommendations.</div>
</div>
""")

# ---------------------------------------------------------
# HOME FOOTER
# ---------------------------------------------------------

if selected_page == "Home":

    st.html("""
<div class="home-footer">
    <span class="home-footer-item"><svg viewBox="0 0 16 16" fill="none"><path d="M3 1.8h6l4 4V14H3V1.8Z" stroke="currentColor" stroke-width="1.2"/><path d="M9 2v4h4M5.5 9h5m-5 2.5h5" stroke="currentColor" stroke-width="1.2" stroke-linecap="round"/></svg>Document Intelligence</span>
    <span class="home-footer-item"><svg viewBox="0 0 16 16" fill="none"><circle cx="4.3" cy="5" r="2.3" stroke="currentColor" stroke-width="1.2"/><circle cx="11.7" cy="11" r="2.3" stroke="currentColor" stroke-width="1.2"/><path d="m6.3 6.4 3.4 3.2" stroke="currentColor" stroke-width="1.2"/></svg>RAG</span>
    <span class="home-footer-item"><svg viewBox="0 0 16 16" fill="none"><path d="M8 1.5 14 4v4c0 3.3-2.5 5.4-6 6.7C4.5 13.4 2 11.3 2 8V4l6-2.5Z" stroke="currentColor" stroke-width="1.2"/><path d="M8 5v3m0 2h.01" stroke="currentColor" stroke-width="1.2" stroke-linecap="round"/></svg>Risk Detection</span>
    <span class="home-footer-item"><svg viewBox="0 0 16 16" fill="none"><path d="M8 1.8 13 3.7v3.5c0 3-2 5.3-5 7-3-1.7-5-4-5-7V3.7l5-1.9Z" stroke="currentColor" stroke-width="1.2"/><path d="m5.5 7.8 1.6 1.6 3.3-3.5" stroke="currentColor" stroke-width="1.2" stroke-linecap="round" stroke-linejoin="round"/></svg>Project Health</span>
    <span class="home-footer-item"><svg viewBox="0 0 16 16" fill="none"><rect x="2" y="3" width="12" height="10" rx="3" stroke="currentColor" stroke-width="1.2"/><circle cx="6" cy="8" r="1" fill="currentColor"/><circle cx="10" cy="8" r="1" fill="currentColor"/><path d="M6 10.5h4" stroke="currentColor" stroke-width="1.2" stroke-linecap="round"/></svg>AI Assistant</span>
</div>
""")

    st.stop()
# ---------------------------------------------------------
# STOP HOME PAGE
# ---------------------------------------------------------

if selected_page == "Home":
    st.stop()



# ---------------------------------------------------------
# DASHBOARD METRICS
# ---------------------------------------------------------

if selected_page == "Dashboard":
    import csv
    import io
    import zipfile

    st.markdown("## AI Project Intelligence Dashboard")
    st.caption("A live view of project health, delivery status, documented risks, and generated analyses.")
    st.markdown("""
    <style>
    .st-key-dashboard_health_dimensions,
    .st-key-dashboard_work_status,
    .st-key-dashboard_risk_visual,
    .st-key-dashboard_risk_summary,
    .st-key-dashboard_blocker_overview {
        border: 1px solid rgba(124, 156, 255, .48);
        border-radius: 16px;
        padding: 12px;
        background: linear-gradient(145deg, rgba(23, 29, 43, .72), rgba(32, 40, 58, .72));
        box-shadow: 0 0 15px rgba(103, 213, 255, .07);
    }
    .st-key-dashboard_work_status { border-color: rgba(126, 214, 165, .48); }
    .st-key-dashboard_risk_visual { border-color: rgba(255, 159, 67, .52); }
    .st-key-dashboard_risk_summary { border-color: rgba(181, 140, 255, .5); }
    .st-key-dashboard_blocker_overview { border-color: rgba(255, 111, 145, .5); }
    </style>
    """, unsafe_allow_html=True)

    processed_documents = st.session_state.get("processed_documents", []) or []
    indexed_chunks = st.session_state.get("indexed_count", 0) or 0
    health_analysis = st.session_state.get("health_analysis", "")
    scope_analysis = st.session_state.get("scope_analysis", "")
    risk_analysis = st.session_state.get("risk_analysis", "")
    blocker_analysis = st.session_state.get("blocker_analysis", "")
    documentation = st.session_state.get("documentation", "")

    health_text = _response_text(health_analysis)
    scope_text = _response_text(scope_analysis).replace("\\n", "\n")
    documentation_text = _response_text(documentation).replace("\\n", "\n")
    health_available = bool(health_text.strip())
    health_score_match = re.search(r"(?im)^\s*Score:\s*(\d{1,3})\s*/\s*100\b", health_text)
    health_score = min(100, int(health_score_match.group(1))) if health_score_match else None
    status_match = re.search(r"(?im)^\s*Status:\s*(.+?)\s*$", health_text)
    health_status = status_match.group(1).strip() if status_match else "Not available"

    analysis_values = [
        ("Scope Analysis", scope_analysis, "document-card"),
        ("Risk Analysis", risk_analysis, "risk-card"),
        ("Blocker Analysis", blocker_analysis, "blocker-card"),
        ("Documentation", documentation, "health-card"),
    ]
    analyses_count = sum(value is not None and bool(str(value).strip()) for _, value, _ in analysis_values)
    dashboard_kpi_slot = st.empty()

    dimensions = [
        "Scope Clarity", "Progress", "Timeline / Delivery",
        "Risk Management", "Blocker Management", "Delivery Confidence",
    ]
    dimension_scores = {}
    for dimension in dimensions:
        match = re.search(rf"(?im)^\s*{re.escape(dimension)}:\s*(.+?)\s*$", health_text)
        score_match = re.search(r"\b(\d{1,3})\s*/\s*100\b", match.group(1)) if match else None
        if score_match:
            dimension_scores[dimension] = min(100, int(score_match.group(1)))

    st.markdown("### Project Health")
    if health_available:
        health_columns = st.columns([2, 1])
        with health_columns[0]:
            st.markdown(
                f"""
                    <div style="padding:20px 22px;border:1px solid #7c9cff66;border-left:4px solid #7c9cff;border-radius:16px;background:radial-gradient(circle at 5% 0%,rgba(124,156,255,.16),transparent 45%),linear-gradient(135deg,#171d2b,#20283a);box-shadow:0 0 16px rgba(124,156,255,.1);">
                    <div style="color:#aeb9ce;font-size:13px;font-weight:600;">Overall Health Score</div>
                    <div style="color:#fff;font-size:30px;font-weight:800;margin:5px 0;">{html.escape(f'{health_score}/100' if health_score is not None else 'Not available')}</div>
                    <div style="color:#9ba8bf;font-size:13px;">Current status: {html.escape(health_status)}</div>
                </div>
                """,
                unsafe_allow_html=True,
            )
            if health_score is not None:
                st.progress(health_score / 100, text="Health score")
            else:
                st.caption("A numeric health score is unavailable.")
        with health_columns[1]:
            st.markdown(
                f"""
                <div style="height:100%;min-height:120px;padding:20px;border:1px solid #7ed6a566;border-left:4px solid #7ed6a5;border-radius:16px;background:linear-gradient(135deg,#171d2b,#20283a);box-shadow:0 0 16px rgba(126,214,165,.08);">
                    <div style="color:#8f9bb2;font-size:12px;text-transform:uppercase;letter-spacing:.08em;">Status</div>
                    <div style="color:#f5f5f5;font-size:22px;font-weight:750;margin-top:12px;">{html.escape(health_status)}</div>
                </div>
                """,
                unsafe_allow_html=True,
            )
    else:
        st.info("Project Health has not been analyzed yet.")

    with st.container(border=True, key="dashboard_health_dimensions"):
        st.markdown("### Health Dimensions")
        st.markdown("<div style='height:2px;width:54px;margin:-6px 0 12px;background:linear-gradient(90deg,#7c9cff,#b58cff);border-radius:4px;'></div>", unsafe_allow_html=True)
        if dimension_scores:
            st.bar_chart(
                {"Dimension": list(dimension_scores), "Score": list(dimension_scores.values())},
                x="Dimension", y="Score", horizontal=True, color="#7c9cff",
            )
        else:
            st.caption("No numeric dimension scores are available in the health analysis.")

    scope_section_labels = (
        "Project Overview", "Project Objective", "Project Scope",
        "Main Modules / Features", "Key Deliverables",
        "Important Requirements", "Important Deadlines / Milestones",
    )
    scope_label_pattern = "|".join(re.escape(label).replace(r"/", r"\s*/\s*") for label in scope_section_labels)
    scope_pipe_pattern = re.compile(
        rf"(?im)(?:^|\|)\s*(?:\d+\s*\|\s*)?(?:\*\*)?(?P<label>{scope_label_pattern})(?:\*\*)?\s*\|"
    )
    scope_matches = list(scope_pipe_pattern.finditer(scope_text))
    dashboard_scope_sections = {}
    if scope_matches:
        for scope_index, scope_match in enumerate(scope_matches):
            scope_end = scope_matches[scope_index + 1].start() if scope_index + 1 < len(scope_matches) else len(scope_text)
            scope_value = scope_text[scope_match.end():scope_end].strip(" \t\r\n|")
            scope_value = re.sub(r"\s*\|\s*", "\n", scope_value).strip()
            if scope_value:
                dashboard_scope_sections[scope_match.group("label").casefold()] = scope_value
    else:
        scope_line_pattern = re.compile(
            rf"(?im)^\s*(?:#{{1,6}}\s*)?(?:\d+[.)]?\s*)?(?P<label>{scope_label_pattern})\s*:?[ \t]*(?P<inline>.*)$"
        )
        scope_lines = scope_text.splitlines()
        scope_heading_matches = [(index, scope_line_pattern.match(line)) for index, line in enumerate(scope_lines)]
        scope_heading_matches = [(index, match) for index, match in scope_heading_matches if match]
        for scope_index, (line_index, scope_match) in enumerate(scope_heading_matches):
            next_line = scope_heading_matches[scope_index + 1][0] if scope_index + 1 < len(scope_heading_matches) else len(scope_lines)
            scope_value = "\n".join(([scope_match.group("inline")] if scope_match.group("inline").strip() else []) + scope_lines[line_index + 1:next_line]).strip()
            if scope_value:
                dashboard_scope_sections[scope_match.group("label").casefold()] = scope_value

    st.markdown("### Scope Summary")
    scope_card_columns = st.columns(2)
    available_scope_sections = [
        (key, label, value) for key, label in (
            ("project overview", "Project Overview"),
            ("project objective", "Project Objective"),
            ("project scope", "Project Scope"),
            ("main modules / features", "Main Modules / Features"),
            ("key deliverables", "Key Deliverables"),
            ("important requirements", "Important Requirements"),
            ("important deadlines / milestones", "Important Deadlines / Milestones"),
        ) if (value := dashboard_scope_sections.get(key))
    ]
    for scope_index, (_, scope_label, scope_value) in enumerate(available_scope_sections):
        with scope_card_columns[scope_index % 2]:
            with st.container(border=True):
                st.markdown(f"#### {html.escape(scope_label)}")
                if scope_label == "Project Overview":
                    overview_fields = (
                        ("Project Name", re.search(r"(?i)\bproject\s*(?:name|title)\s*[:：]\s*(.+?)(?=\s*(?:;|\n|\|)\s*(?:project\s*)?(?:objective|scope)\s*[:：]|$)", scope_value)),
                        ("Project Objective", re.search(r"(?i)\b(?:project\s*)?objective\s*[:：]\s*(.+?)(?=\s*(?:;|\n|\|)\s*(?:project\s*)?(?:name|scope)\s*[:：]|$)", scope_value)),
                    )
                    overview_remaining = scope_value
                    for field_label, field_match in overview_fields:
                        if field_match:
                            field_value = field_match.group(1).strip(" *|;")
                            st.markdown(f"**{field_label}:** {field_value}")
                            overview_remaining = overview_remaining.replace(field_match.group(0), "", 1)
                    overview_remaining = re.sub(r"(?i)\b(?:project\s*)?(?:name|title|objective)\s*[:：]", "", overview_remaining).strip(" *|;\n")
                    if overview_remaining:
                        st.markdown(overview_remaining)
                else:
                    scope_items = [item.strip(" -*•") for item in re.split(r"[\n;•]+", scope_value) if item.strip(" -*•")]
                    if len(scope_items) == 1 and scope_label in ("Project Scope", "Main Modules / Features"):
                        scope_items = [item.strip() for item in scope_items[0].split(",") if item.strip()]
                    if len(scope_items) == 1 and scope_label == "Important Requirements":
                        scope_items = [item.strip() for item in re.split(r"(?<=[.!?])\s+(?=[A-Z0-9])", scope_items[0]) if item.strip()]
                    for scope_item in scope_items:
                        st.markdown(f"- {scope_item}")
    if not available_scope_sections:
        st.caption("No structured scope sections are available in the current analysis.")

    work_status_records = {}
    work_status_categories = ("Completed", "In Progress", "Blocked", "Not Started")
    work_status_reliable = True
    work_tables_found = 0
    for document in processed_documents:
        is_csv = str(document.get("file_type", "")).upper() == "CSV" or str(document.get("filename", "")).lower().endswith(".csv")
        if not is_csv:
            continue
        lines = str(document.get("text", "")).splitlines()
        header_index = next((index for index, line in enumerate(lines) if re.search(r"\bTask[_ ]?ID\b", line, re.IGNORECASE) and re.search(r"\bStatus\b", line, re.IGNORECASE)), None)
        if header_index is None:
            continue
        work_tables_found += 1
        rows_found = 0
        for line in lines[header_index + 1:]:
            if not line.strip():
                continue
            row_match = re.match(r"^\s*(?P<task_id>\S+)\s+.+?\s+(?P<status>Completed|In Progress|Blocked|Not Started)(?:\s+\d{4}-\d{2}-\d{2})?\s*$", line, re.IGNORECASE)
            if not row_match:
                work_status_reliable = False
                break
            task_id = row_match.group("task_id")
            status_value = next(category for category in work_status_categories if category.casefold() == row_match.group("status").casefold())
            if task_id in work_status_records and work_status_records[task_id] != status_value:
                work_status_reliable = False
                break
            work_status_records[task_id] = status_value
            rows_found += 1
        if rows_found == 0:
            work_status_reliable = False

    work_status_counts = None
    if work_tables_found and work_status_records and work_status_reliable:
        work_status_counts = {status: sum(value == status for value in work_status_records.values()) for status in work_status_categories}
    work_status_total = sum(work_status_counts.values()) if work_status_counts else 0
    work_progress = round(work_status_counts.get("Completed", 0) / work_status_total * 100) if work_status_total else None

    risk_records = []
    risk_text = _response_text(risk_analysis).replace("\\n", "\n")
    risk_register_match = re.search(
        r"(?ims)^\s*#\s*(?:2\.\s*)?RISK REGISTER\s*$\s*(.*?)(?=^\s*#\s*(?:3\.\s*)?ACTION ITEMS\b|\Z)",
        documentation_text,
    )
    risk_register_text = risk_register_match.group(1) if risk_register_match else ""
    risk_text = "\n\n".join(part for part in (risk_text, risk_register_text) if part.strip())
    risk_blocks = re.split(r"(?im)(?=^\s*(?:[-*•]\s*)?(?:\d+[.)]\s*)?(?:\*\*)?Risk(?:\s+\d+)?(?:\*\*)?\s*[:\-]\s*\S)", risk_text)
    for block in risk_blocks:
        name_match = re.search(r"(?im)^\s*(?:[-*•]\s*)?(?:\d+[.)]\s*)?(?:\*\*)?Risk(?:\s+\d+)?(?:\*\*)?\s*[:\-]\s*(?:\*\*)?(.+?)(?:\*\*)?\s*$", block)
        if not name_match:
            continue
        record = {"Risk": name_match.group(1).strip().strip("*")}
        for field in ("Severity", "Priority", "Impact", "Potential Project Impact", "Probability", "Likelihood", "Status", "Owner", "Reason", "Recommended Action"):
            field_match = re.search(rf"(?im)^\s*(?:[-*•]\s*)?(?:\*\*)?{field}(?:\*\*)?\s*[:\-]\s*(.+?)\s*$", block)
            if field_match:
                value = field_match.group(1).strip().strip("*")
                if field == "Potential Project Impact" and not record.get("Impact"):
                    record["Impact"] = value
                elif field in ("Impact", "Probability", "Likelihood", "Severity", "Priority"):
                    record[field] = value
                    category_match = re.fullmatch(r"(?i)(High|Medium|Low)\.?", value)
                    if category_match:
                        record[field] = category_match.group(1).title()
                elif value and not value.casefold().startswith(("not specified", "not found")):
                    record[field] = value
        risk_records.append(record)

    risk_field_line = re.compile(
        r"(?i)^\s*(?:[-*•]\s*)?(?:\*\*)?(?:Risk|Severity|Priority|Impact|Probability|"
        r"Likelihood|Status|Owner|Reason|Potential Project Impact|Recommended Action)"
        r"(?:\*\*)?\s*[:\-]"
    )
    risk_evidence = re.compile(
        r"(?i)\b(?:risk|delay(?:ed)?|blocked|blocker|pending|not finalized|"
        r"unfinished|limited(?:\s+[\w-]+){0,3}\s+availability|"
        r"availability limitation|dependency|"
        r"schedule slip|schedule delay|behind schedule|at risk)\b"
    )
    known_risks = {record["Risk"].casefold() for record in risk_records}
    for source_text, is_risk_register in ((risk_register_text, True), (risk_text, False)):
        for line in source_text.splitlines():
            if risk_field_line.match(line):
                continue
            is_list_item = bool(re.match(r"^\s*(?:[-*•]|\d+[.)])\s+", line))
            entry_match = re.match(r"^\s*(?:[-*•]\s*|\d+[.)]\s*)?(.+?)\s*$", line)
            if not entry_match:
                continue
            risk_name = entry_match.group(1).replace("**", "").strip()
            risk_name = re.sub(r"(?i)^risk\s*[:\-]\s*", "", risk_name).strip()
            if (
                not risk_name
                or len(risk_name) > 500
                or risk_name.startswith(("#", "="))
                or risk_name.casefold().startswith(("not found", "no documented risks", "no risks identified"))
                or (is_risk_register and not is_list_item and not risk_evidence.search(risk_name))
                or not (is_risk_register or risk_evidence.search(risk_name))
                or risk_name.casefold() in known_risks
            ):
                continue
            risk_records.append({"Risk": risk_name})
            known_risks.add(risk_name.casefold())

    def risk_category(value):
        match = re.match(r"(?i)^\s*(Low|Medium|High)\b", str(value or ""))
        return match.group(1).title() if match else None

    risk_severity_counts = {
        severity: sum(risk_category(record.get("Severity") or record.get("Priority")) == severity for record in risk_records)
        for severity in ("High", "Medium", "Low")
    }
    if not any(risk_severity_counts.values()):
        risk_severity_counts = None
    risk_points = []
    category_scale = {"Low": 1, "Medium": 2, "High": 3}

    risk_impact_levels = [risk_category(record.get("Impact")) for record in risk_records]
    risk_severity_levels = [risk_category(record.get("Severity") or record.get("Priority")) for record in risk_records]
    risk_distribution_counts = None
    if risk_records and any(risk_impact_levels):
        levels = [level or "Not specified" for level in risk_impact_levels]
        risk_distribution_counts = {level: sum(value == level for value in levels) for level in dict.fromkeys(levels)}
    elif risk_records and any(risk_severity_levels):
        levels = [level or "Not specified" for level in risk_severity_levels]
        risk_distribution_counts = {level: sum(value == level for value in levels) for level in dict.fromkeys(levels)}

    for record in risk_records:
        impact = risk_category(record.get("Impact"))
        probability = risk_category(record.get("Probability", record.get("Likelihood", "")))
        if impact in category_scale and probability in category_scale:
            risk_points.append({
                "Risk": record["Risk"], "Impact": record.get("Impact", ""),
                "Probability": record.get("Probability", record.get("Likelihood", "")),
                "Status": record.get("Status", ""),
                "Impact Score": category_scale[impact],
                "Probability Score": category_scale[probability],
            })

    blocker_text = _response_text(blocker_analysis).replace("\\n", "\n")
    blocker_blocks = re.split(r"(?im)(?=^###\s+Blocker\s+\d+\b)", blocker_text)
    blocker_records = []
    for block in blocker_blocks:
        if not re.search(r"(?im)^###\s+Blocker\s+\d+\b", block):
            continue
        name_match = re.search(r"(?im)^\s*-\s*\*\*Blocker:\*\*\s*(.+?)\s*$", block)
        if not name_match or re.search(r"(?i)^(no blocker identified|unable to analyze blockers)", name_match.group(1).strip()):
            continue
        record = {"Blocker": name_match.group(1).strip()}
        for field in ("Description", "Priority", "Owner", "Required Action", "Status"):
            field_match = re.search(rf"(?im)^\s*-\s*\*\*{field}:\*\*\s*(.+?)\s*$", block)
            if field_match:
                value = field_match.group(1).strip()
                if not value.casefold().startswith(("not specified", "not found")):
                    record[field] = value.title() if field == "Priority" else value
        blocker_records.append(record)

    blocker_priority_counts = {
        priority: sum(record.get("Priority") == priority for record in blocker_records)
        for priority in ("High", "Medium", "Low")
    }
    if not any(blocker_priority_counts.values()):
        blocker_priority_counts = None
    blocker_status_counts = None
    if blocker_records and all(record.get("Status") for record in blocker_records):
        blocker_status_counts = {}
        for record in blocker_records:
            blocker_status_counts[record["Status"]] = blocker_status_counts.get(record["Status"], 0) + 1

    action_section_match = re.search(
        r"(?ims)^\s*#*\s*(?:3\.\s*)?ACTION ITEMS\s*$\s*(.*?)(?=^\s*#*\s*(?:4\.\s*)?PROJECT SUMMARY\b|\Z)",
        documentation_text,
    )
    action_records = []
    if action_section_match:
        for action_line in action_section_match.group(1).splitlines():
            action_match = re.match(r"^\s*(?:[-*•]|\d+[.)])\s+(.+?)\s*$", action_line)
            if not action_match:
                continue
            action_text = action_match.group(1).strip()
            if not action_text or action_text.casefold().startswith(("not found", "no documented actions")):
                continue
            action_record = {"Action": action_text}
            for field in ("Priority", "Status"):
                field_match = re.search(rf"(?i)\b{field}\s*:\s*([^;,|]+)", action_text)
                if field_match and not field_match.group(1).strip().casefold().startswith(("not specified", "not found")):
                    action_record[field] = field_match.group(1).strip()
            action_records.append(action_record)
    if not action_records:
        for record in blocker_records:
            action_text = record.get("Required Action", "").strip()
            if action_text and action_text.casefold() not in {item["Action"].casefold() for item in action_records}:
                action_records.append({
                    "Action": action_text,
                    **({"Priority": record["Priority"]} if record.get("Priority") else {}),
                    **({"Status": record["Status"]} if record.get("Status") else {}),
                })
    if not action_records:
        for record in risk_records:
            for action_line in re.split(r"[\n;]+", record.get("Recommended Action", "")):
                action_text = action_line.strip(" -*•")
                if action_text and action_text.casefold() not in {item["Action"].casefold() for item in action_records}:
                    action_records.append({"Action": action_text})
    if not action_records:
        health_actions_match = re.search(
            r"(?ims)^\s*#\s*IMMEDIATE ACTIONS REQUIRED\s*$\s*(.*?)(?=^\s*#\s*[^\n]+|\Z)",
            health_text,
        )
        if health_actions_match:
            for action_line in health_actions_match.group(1).splitlines():
                action_match = re.match(r"^\s*(?:[-*•]|\d+[.)])\s+(.+?)\s*$", action_line)
                if action_match and not action_match.group(1).casefold().startswith(("not found", "insufficient information")):
                    action_records.append({"Action": action_match.group(1).strip()})

    dashboard_metrics = [
        ("Project Health Score", f"{health_score}/100" if health_score is not None else ("Score unavailable" if health_available else "Not analyzed"), "#7ed6a5", '<svg viewBox="0 0 24 24"><path d="M20.8 8.7c0 5.1-8.8 11-8.8 11s-8.8-5.9-8.8-11A4.7 4.7 0 0 1 12 6.4a4.7 4.7 0 0 1 8.8 2.3Z"/><path d="M5 12h4l2-3 3 6 2-3h3"/></svg>'),
        ("Total Risks", len(risk_records) if risk_records else "—", "#ff9f43", '<svg viewBox="0 0 24 24"><path d="m12 3 10 18H2L12 3Z"/><path d="M12 9v5m0 3h.01"/></svg>'),
        ("Total Blockers", len(blocker_records) if blocker_records else "—", "#ff6f91", '<svg viewBox="0 0 24 24"><path d="M8 6h13M8 12h13M8 18h13M3 6h.01M3 12h.01M3 18h.01"/></svg>'),
        ("Total Action Items", len(action_records) if action_records else "—", "#b58cff", '<svg viewBox="0 0 24 24"><path d="M8 6h13M8 12h13M8 18h13M3 6h.01M3 12h.01M3 18h.01"/></svg>'),
    ]
    if work_progress is not None:
        dashboard_metrics.append(("Project Progress", f"{work_progress}%", "#67d5ff", '<svg viewBox="0 0 24 24"><path d="M4 19V5M4 19h17M8 16v-4m5 4V8m5 8V5"/></svg>'))
    with dashboard_kpi_slot.container():
        metric_columns = st.columns(len(dashboard_metrics))
        for column, (title, value, accent, icon_svg) in zip(metric_columns, dashboard_metrics):
            with column:
                st.markdown(
                    f"<div class='metric-card' style='min-height:112px;text-align:left;background:linear-gradient(145deg,#171d2b,#20283a);border:1px solid {accent}66;border-left:4px solid {accent};border-radius:16px;padding:16px;box-shadow:0 8px 24px rgba(0,0,0,.16),0 0 14px {accent}15;'><div style='display:flex;align-items:center;justify-content:space-between;gap:10px;'><div><div class='metric-title'>{html.escape(title)}</div><div class='metric-value'>{html.escape(str(value))}</div></div><span style='width:38px;height:38px;display:grid;place-items:center;border-radius:11px;background:{accent}1c;border:1px solid {accent}55;color:{accent};'>{icon_svg.replace('<svg ', '<svg width="22" height="22" fill="none" stroke="currentColor" stroke-width="1.7" stroke-linecap="round" stroke-linejoin="round" ')}</span></div></div>",
                    unsafe_allow_html=True,
                )

    st.markdown("### Project Work Status")
    with st.container(border=True, key="dashboard_work_status"):
        st.markdown("<div style='height:2px;width:54px;margin:-6px 0 12px;background:linear-gradient(90deg,#7ed6a5,#67d5ff);border-radius:4px;'></div>", unsafe_allow_html=True)
        if work_status_counts:
            st.bar_chart({"Status": list(work_status_counts), "Tasks": list(work_status_counts.values())}, x="Status", y="Tasks", horizontal=True, color="#7ed6a5")
        else:
            st.caption("Reliable structured task status data is unavailable.")

    st.markdown("### Risk Overview")
    risk_visual_column, risk_summary_column = st.columns([1.15, 1.25])
    with risk_visual_column:
        with st.container(border=True, key="dashboard_risk_visual"):
            st.markdown("#### Risk Matrix / Distribution")
            st.markdown("<div style='height:2px;width:54px;margin:-6px 0 12px;background:linear-gradient(90deg,#ff9f43,#ff6f91);border-radius:4px;'></div>", unsafe_allow_html=True)
            if risk_points and len(risk_points) == len(risk_records) and hasattr(st, "scatter_chart"):
                st.scatter_chart(risk_points, x="Impact Score", y="Probability Score", color="Risk")
                st.caption("Hover chart points for risk details. Plot coordinates map documented Low, Medium, and High values to 1, 2, and 3; original text values remain in the table.")
            elif risk_distribution_counts:
                distribution_label = "Impact" if any(risk_impact_levels) else "Severity"
                st.bar_chart({distribution_label: list(risk_distribution_counts), "Risks": list(risk_distribution_counts.values())}, x=distribution_label, y="Risks", horizontal=True, color="#ff9f43")
            elif risk_records:
                st.bar_chart({"Category": ["Documented risks"], "Count": [len(risk_records)]}, x="Category", y="Count", color="#ff9f43")
            else:
                st.caption("No documented risk entries are available in Risk Analysis or the generated Risk Register.")
            if risk_severity_counts:
                st.markdown("#### Risks by Severity")
                st.bar_chart(
                    {"Severity": list(risk_severity_counts), "Risks": list(risk_severity_counts.values())},
                    x="Severity", y="Risks", color="#ff9f43",
                )
                if sum(risk_severity_counts.values()) < len(risk_records):
                    st.caption("Only risks with documented severity or priority are included in this chart.")
    with risk_summary_column:
        with st.container(border=True, key="dashboard_risk_summary"):
            st.markdown("#### Risk Summary")
            st.markdown("<div style='height:2px;width:54px;margin:-6px 0 12px;background:linear-gradient(90deg,#b58cff,#67d5ff);border-radius:4px;'></div>", unsafe_allow_html=True)
            if risk_distribution_counts:
                risk_summary_columns = st.columns(3)
                for column, level in zip(risk_summary_columns, ("High", "Medium", "Low")):
                    with column:
                        st.metric(level, risk_distribution_counts.get(level, 0))
                if "Not specified" in risk_distribution_counts:
                    st.caption(f"Not specified: {risk_distribution_counts['Not specified']}")
            if risk_records:
                risk_filtered_records = list(risk_records)
                risk_severity_options = sorted({record.get("Severity") or record.get("Priority") for record in risk_records if record.get("Severity") or record.get("Priority")})
                risk_status_options = sorted({record["Status"] for record in risk_records if record.get("Status")})
                risk_filter_columns = st.columns(2)
                selected_risk_severities = []
                selected_risk_statuses = []
                if risk_severity_options:
                    with risk_filter_columns[0]:
                        selected_risk_severities = st.multiselect("Filter severity", risk_severity_options, key="dashboard_risk_severity_filter")
                if risk_status_options:
                    with risk_filter_columns[1]:
                        selected_risk_statuses = st.multiselect("Filter status", risk_status_options, key="dashboard_risk_status_filter")
                risk_filtered_records = [
                    record for record in risk_records
                    if (not selected_risk_severities or (record.get("Severity") or record.get("Priority")) in selected_risk_severities)
                    and (not selected_risk_statuses or record.get("Status") in selected_risk_statuses)
                ]
                risk_display_fields = [
                    ("Risk Name", "Risk"), ("Severity", "Severity"), ("Priority", "Priority"),
                    ("Impact", "Impact"), ("Probability", "Probability"),
                    ("Status", "Status"), ("Owner", "Owner"),
                ]
                risk_display_fields = [
                    (label, key) for label, key in risk_display_fields
                    if key == "Risk" or any(record.get(key) or (key == "Probability" and record.get("Likelihood")) for record in risk_records)
                ]
                risk_table = []
                for record in risk_filtered_records:
                    row = {}
                    for label, key in risk_display_fields:
                        row[label] = record.get("Probability", record.get("Likelihood", "")) if key == "Probability" else record.get(key, "")
                    risk_table.append(row)
                st.dataframe(
                    risk_table,
                    hide_index=True,
                    use_container_width=True,
                    height=min(300, 44 + 36 * len(risk_table)),
                    column_config={
                        label: st.column_config.TextColumn(label, width="small" if label != "Risk Name" else "medium")
                        for label, _ in risk_display_fields
                    },
                )
                if risk_filtered_records:
                    with st.expander("Inspect Risk Details"):
                        selected_risk_index = st.selectbox(
                            "Select a risk",
                            range(len(risk_filtered_records)),
                            format_func=lambda index: risk_filtered_records[index]["Risk"],
                            key="dashboard_risk_detail_select",
                        )
                        selected_risk = risk_filtered_records[selected_risk_index]
                        for field in ("Severity", "Priority", "Impact", "Probability", "Status", "Owner", "Reason", "Recommended Action"):
                            if selected_risk.get(field):
                                st.markdown(f"**{field}:** {selected_risk[field]}")
            else:
                st.caption("No documented risks are available to display.")

    st.markdown("### Blocker Overview")
    with st.container(border=True, key="dashboard_blocker_overview"):
        st.markdown("<div style='height:2px;width:54px;margin:-6px 0 12px;background:linear-gradient(90deg,#ff6f91,#ff9f43);border-radius:4px;'></div>", unsafe_allow_html=True)
        if blocker_priority_counts:
            st.bar_chart({"Priority": list(blocker_priority_counts), "Blockers": list(blocker_priority_counts.values())}, x="Priority", y="Blockers", horizontal=True, color="#ff6f91")
        elif blocker_status_counts:
            st.bar_chart({"Status": list(blocker_status_counts), "Blockers": list(blocker_status_counts.values())}, x="Status", y="Blockers", horizontal=True, color="#ff6f91")
        elif blocker_records:
            st.caption(f"{len(blocker_records)} documented blocker(s); priority and status are not reliably available.")
        else:
            st.caption("No reliably structured blocker entries are available.")
    if blocker_records:
        blocker_priority_options = sorted({record["Priority"] for record in blocker_records if record.get("Priority")})
        blocker_status_options = sorted({record["Status"] for record in blocker_records if record.get("Status")})
        blocker_filter_columns = st.columns(2)
        selected_blocker_priorities = []
        selected_blocker_statuses = []
        if blocker_priority_options:
            with blocker_filter_columns[0]:
                selected_blocker_priorities = st.multiselect("Filter blocker priority", blocker_priority_options, key="dashboard_blocker_priority_filter")
        if blocker_status_options:
            with blocker_filter_columns[1]:
                selected_blocker_statuses = st.multiselect("Filter blocker status", blocker_status_options, key="dashboard_blocker_status_filter")
        filtered_blockers = [
            record for record in blocker_records
            if (not selected_blocker_priorities or record.get("Priority") in selected_blocker_priorities)
            and (not selected_blocker_statuses or record.get("Status") in selected_blocker_statuses)
        ]
        with st.expander("Inspect Blocker Details"):
            for index, record in enumerate(filtered_blockers, start=1):
                st.markdown(f"**{index}. {record['Blocker']}**")
                for field in ("Description", "Priority", "Owner", "Required Action", "Status"):
                    if record.get(field):
                        st.markdown(f"- **{field}:** {record[field]}")

    st.markdown("### Action Items")
    if action_records:
        action_priority_options = sorted({record["Priority"] for record in action_records if record.get("Priority")})
        action_status_options = sorted({record["Status"] for record in action_records if record.get("Status")})
        action_filter_columns = st.columns(2)
        selected_action_priorities = []
        selected_action_statuses = []
        if action_priority_options:
            with action_filter_columns[0]:
                selected_action_priorities = st.multiselect("Filter action priority", action_priority_options, key="dashboard_action_priority_filter")
        if action_status_options:
            with action_filter_columns[1]:
                selected_action_statuses = st.multiselect("Filter action status", action_status_options, key="dashboard_action_status_filter")
        filtered_actions = [
            record for record in action_records
            if (not selected_action_priorities or record.get("Priority") in selected_action_priorities)
            and (not selected_action_statuses or record.get("Status") in selected_action_statuses)
        ]
        st.dataframe(filtered_actions, hide_index=True, use_container_width=True)
    else:
        st.caption("No action items are available in the current documentation or blocker analysis.")

    coverage = [
        ("Scope Analysis", scope_analysis, "#67d5ff", '<svg viewBox="0 0 24 24"><circle cx="12" cy="12" r="8"/><circle cx="12" cy="12" r="3"/><path d="M12 2v3m0 14v3M2 12h3m14 0h3"/></svg>'),
        ("Risk Analysis", risk_analysis, "#ff9f43", '<svg viewBox="0 0 24 24"><path d="m12 3 10 18H2L12 3Z"/><path d="M12 9v5m0 3h.01"/></svg>'),
        ("Blocker Analysis", blocker_analysis, "#ff6f91", '<svg viewBox="0 0 24 24"><path d="M8 6h13M8 12h13M8 18h13M3 6h.01M3 12h.01M3 18h.01"/></svg>'),
        ("Health Analysis", health_analysis, "#7ed6a5", '<svg viewBox="0 0 24 24"><path d="M20.8 8.7c0 5.1-8.8 11-8.8 11s-8.8-5.9-8.8-11A4.7 4.7 0 0 1 12 6.4a4.7 4.7 0 0 1 8.8 2.3Z"/><path d="M5 12h4l2-3 3 6 2-3h3"/></svg>'),
        ("Documentation", documentation, "#b58cff", '<svg viewBox="0 0 24 24"><path d="M6 3h8l4 4v14H6zM14 3v5h5M9 13h6M9 17h6"/></svg>'),
    ]
    st.markdown("### Analysis Coverage")
    coverage_columns = st.columns(5)
    for column, (title, value, accent, icon_svg) in zip(coverage_columns, coverage):
        available = value is not None and bool(str(value).strip())
        with column:
            st.markdown(
                f"<div class='metric-card' style='min-height:92px;text-align:left;background:linear-gradient(145deg,#171d2b,#20283a);border:1px solid {accent}66;border-left:3px solid {accent};border-radius:14px;padding:13px;box-shadow:0 0 12px {accent}16;'><div style='display:flex;align-items:center;gap:9px;margin-bottom:7px;'><span style='width:19px;height:19px;color:{accent};display:inline-flex;'>{icon_svg.replace('<svg ', '<svg width=\"19\" height=\"19\" fill=\"none\" stroke=\"currentColor\" stroke-width=\"1.7\" stroke-linecap=\"round\" stroke-linejoin=\"round\" ')}</span><div class='metric-title'>{html.escape(title)}</div></div><div style='font-size:14px;font-weight:700;color:{'#7ed6a5' if available else '#9aa3b5'};'>{'Available' if available else 'Unavailable'}</div></div>",
                unsafe_allow_html=True,
            )

    # The report/export data mirrors the values shown above and never triggers new analyses.
    health_dimension_rows = [[name, score] for name, score in dimension_scores.items()]
    work_status_rows = [[name, count] for name, count in (work_status_counts or {}).items()]
    risk_rows = [[record.get("Risk", ""), record.get("Severity", ""), record.get("Impact", ""), record.get("Probability", record.get("Likelihood", "")), record.get("Status", ""), record.get("Owner", "")] for record in risk_records]
    blocker_rows = [[record.get("Blocker", ""), record.get("Priority", ""), record.get("Status", "")] for record in blocker_records]
    coverage_rows = [[name, "Available" if value is not None and bool(str(value).strip()) else "Not Generated"] for name, value, _, _ in coverage]
    summary_rows = [
        ["Documents", len(processed_documents)], ["Indexed Chunks", indexed_chunks],
        ["Project Health Score", f"{health_score}/100" if health_score is not None else ""],
        ["Project Health Status", health_status if health_available else "Not analyzed"],
        ["Analyses Available", f"{analyses_count} / 4"],
    ]
    report_sheets = [
        ("Dashboard Summary", [ ["Metric", "Value"] ] + summary_rows),
        ("Health Dimensions", [["Dimension", "Score / 100"]] + health_dimension_rows),
        ("Work Status", [["Status", "Count"]] + work_status_rows),
        ("Risks", [["Risk", "Severity", "Impact", "Probability", "Status", "Owner"]] + risk_rows),
        ("Blockers", [["Blocker", "Priority", "Status"]] + blocker_rows),
        ("Analysis Coverage", [["Analysis", "Status"]] + coverage_rows),
    ]

    dashboard_csv = io.StringIO(newline="")
    csv_writer = csv.writer(dashboard_csv)
    csv_writer.writerow(["Section", "Item", "Value", "Details"])
    for name, value in summary_rows:
        csv_writer.writerow(["Dashboard Summary", name, value, ""])
    for name, value in health_dimension_rows:
        csv_writer.writerow(["Health Dimensions", name, value, ""])
    for name, value in work_status_rows:
        csv_writer.writerow(["Work Status", name, value, ""])
    for risk in risk_rows:
        csv_writer.writerow(["Risks", risk[0], risk[1], "; ".join(f"{label}: {risk[index]}" for index, label in ((2, "Impact"), (3, "Probability"), (4, "Status"), (5, "Owner")) if risk[index])])
    for blocker in blocker_rows:
        csv_writer.writerow(["Blockers", blocker[0], blocker[1], f"Status: {blocker[2]}" if blocker[2] else ""])
    for name, value in coverage_rows:
        csv_writer.writerow(["Analysis Coverage", name, value, ""])
    csv_bytes = dashboard_csv.getvalue().encode("utf-8-sig")

    def excel_column(number):
        result = ""
        while number:
            number, remainder = divmod(number - 1, 26)
            result = chr(65 + remainder) + result
        return result

    excel_buffer = io.BytesIO()
    with zipfile.ZipFile(excel_buffer, "w", zipfile.ZIP_DEFLATED) as workbook:
        sheet_names = [name for name, _ in report_sheets]
        content_types = ['<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>']
        workbook_sheets = []
        rel_sheets = []
        for sheet_index, (sheet_name, rows) in enumerate(report_sheets, start=1):
            content_types.append(f'<Override PartName="/xl/worksheets/sheet{sheet_index}.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>')
            workbook_sheets.append(f'<sheet name="{html.escape(sheet_name, quote=True)}" sheetId="{sheet_index}" r:id="rId{sheet_index}"/>')
            rel_sheets.append(f'<Relationship Id="rId{sheet_index}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet{sheet_index}.xml"/>')
            sheet_rows = []
            for row_index, row in enumerate(rows, start=1):
                cells = []
                for column_index, value in enumerate(row, start=1):
                    reference = f"{excel_column(column_index)}{row_index}"
                    if isinstance(value, (int, float)) and not isinstance(value, bool):
                        cells.append(f'<c r="{reference}" t="n"><v>{value}</v></c>')
                    else:
                        safe_value = re.sub(r"[\x00-\x08\x0B\x0C\x0E-\x1F]", "", str(value))
                        cells.append(f'<c r="{reference}" t="inlineStr"><is><t xml:space="preserve">{html.escape(safe_value)}</t></is></c>')
                sheet_rows.append(f'<row r="{row_index}">{"".join(cells)}</row>')
            workbook.writestr(f"xl/worksheets/sheet{sheet_index}.xml", '<?xml version="1.0" encoding="UTF-8"?><worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>' + "".join(sheet_rows) + "</sheetData></worksheet>")
        content_types.append("</Types>")
        workbook.writestr("[Content_Types].xml", "".join(content_types))
        workbook.writestr("_rels/.rels", '<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/></Relationships>')
        workbook.writestr("xl/workbook.xml", '<?xml version="1.0" encoding="UTF-8"?><workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets>' + "".join(workbook_sheets) + "</sheets></workbook>")
        workbook.writestr("xl/_rels/workbook.xml.rels", '<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">' + "".join(rel_sheets) + "</Relationships>")
    excel_bytes = excel_buffer.getvalue()

    pdf_lines = [("AI Project Intelligence Dashboard", "title"), ("Dashboard Summary", "heading")]
    pdf_lines.extend((f"{name}: {value}", "body") for name, value in summary_rows)
    pdf_lines.append(("Health Dimensions", "heading"))
    if health_dimension_rows:
        pdf_lines.extend((f"{name}: {score}/100", "body") for name, score in health_dimension_rows)
    else:
        pdf_lines.append(("Data unavailable", "body"))
    pdf_lines.append(("Project Work Status", "heading"))
    if work_status_rows:
        pdf_lines.extend((f"{name}: {count}", "body") for name, count in work_status_rows)
    else:
        pdf_lines.append(("Data unavailable", "body"))
    pdf_lines.append(("Risks", "heading"))
    if risk_rows:
        pdf_lines.extend((f"{risk[0]} | Severity: {risk[1]} | Impact: {risk[2]} | Probability: {risk[3]} | Status: {risk[4]} | Owner: {risk[5]}", "body") for risk in risk_rows)
    else:
        pdf_lines.append(("Data unavailable", "body"))
    pdf_lines.append(("Blockers", "heading"))
    if blocker_rows:
        pdf_lines.extend((f"{blocker[0]} | Priority: {blocker[1]} | Status: {blocker[2]}", "body") for blocker in blocker_rows)
    else:
        pdf_lines.append(("Data unavailable", "body"))
    pdf_lines.append(("Analysis Coverage", "heading"))
    pdf_lines.extend(((f"{name}: {status}", "body") for name, status in coverage_rows))

    wrapped_pdf_lines = []
    for line_text, style in pdf_lines:
        plain = str(line_text).replace("**", "").replace("•", "-").replace("–", "-").replace("—", "-")
        plain = plain.encode("ascii", "ignore").decode("ascii")
        while len(plain) > 92:
            split_at = plain.rfind(" ", 0, 92)
            if split_at < 1:
                split_at = 92
            wrapped_pdf_lines.append((plain[:split_at], style))
            plain = plain[split_at:].lstrip()
        wrapped_pdf_lines.append((plain, style))
    pdf_pages, page_commands, y_position = [], [], 790
    for line_text, line_style in wrapped_pdf_lines:
        if y_position < 55:
            pdf_pages.append("\n".join(page_commands).encode("ascii"))
            page_commands, y_position = [], 790
        font_name, font_size, line_gap = {"title": ("F2", 18, 32), "heading": ("F2", 13, 24), "body": ("F1", 9, 14)}[line_style]
        escaped_text = line_text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        page_commands.append(f"BT /{font_name} {font_size} Tf 50 {y_position} Td ({escaped_text}) Tj ET")
        y_position -= line_gap
    if page_commands:
        pdf_pages.append("\n".join(page_commands).encode("ascii"))
    pdf_objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b"", b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>", b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold >>"]
    page_ids = []
    for page_index, page_content in enumerate(pdf_pages):
        page_id = 5 + page_index * 2
        stream_id = page_id + 1
        page_ids.append(page_id)
        pdf_objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 842] /Resources << /Font << /F1 3 0 R /F2 4 0 R >> >> /Contents {stream_id} 0 R >>".encode("ascii"))
        pdf_objects.append(f"<< /Length {len(page_content)} >>\nstream\n".encode("ascii") + page_content + b"\nendstream")
    pdf_objects[1] = f"<< /Type /Pages /Kids [{' '.join(f'{page_id} 0 R' for page_id in page_ids)}] /Count {len(page_ids)} >>".encode("ascii")
    pdf_buffer = bytearray(b"%PDF-1.4\n")
    object_offsets = [0]
    for object_id, pdf_object in enumerate(pdf_objects, start=1):
        object_offsets.append(len(pdf_buffer))
        pdf_buffer.extend(f"{object_id} 0 obj\n".encode("ascii"))
        pdf_buffer.extend(pdf_object + b"\nendobj\n")
    xref_offset = len(pdf_buffer)
    pdf_buffer.extend(f"xref\n0 {len(object_offsets)}\n".encode("ascii"))
    pdf_buffer.extend(b"0000000000 65535 f \n")
    for offset in object_offsets[1:]:
        pdf_buffer.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    pdf_buffer.extend(f"trailer\n<< /Size {len(object_offsets)} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF".encode("ascii"))
    pdf_bytes = bytes(pdf_buffer)

    st.markdown("### Download Dashboard")
    st.caption("Downloads use the current dashboard/session data; no analysis is regenerated.")
    download_columns = st.columns(3)
    with download_columns[0]:
        with st.container(border=True):
            st.download_button("Download Dashboard PDF", pdf_bytes, "project_dashboard.pdf", "application/pdf", use_container_width=True, key="dashboard_download_pdf")
    with download_columns[1]:
        with st.container(border=True):
            st.download_button("Download Dashboard Excel", excel_bytes, "project_dashboard.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", use_container_width=True, key="dashboard_download_excel")
    with download_columns[2]:
        with st.container(border=True):
            st.download_button("Download Dashboard CSV", csv_bytes, "project_dashboard.csv", "text/csv", use_container_width=True, key="dashboard_download_csv")

    st.divider()
    dashboard_navigation_columns = st.columns([1, 2, 1])
    with dashboard_navigation_columns[0]:
        if st.button("← Back to Home", key="dashboard_back_home"):
            st.session_state.selected_page = "Home"
            st.rerun()
    with dashboard_navigation_columns[2]:
        if st.button("Go to AI Assistant →", key="dashboard_next_ai_assistant", use_container_width=True):
            st.session_state.selected_page = "AI Assistant"
            st.rerun()


# ---------------------------------------------------------
# DOCUMENTS PAGE
# ---------------------------------------------------------

if selected_page == "Documents":

    st.markdown("""
    <div class="module-card" style="--accent:#67d5ff;--progress:45%;">
      <div class="module-icon">📁</div>
      <div class="module-title">Upload Project Documents</div>
      <span class="module-badge">Knowledge Base</span>
      <div class="module-description">
        Upload multiple project artifacts to build the project knowledge base.
      </div>
      <div class="module-line"><span></span></div>
      <div class="module-caption">PDF • DOCX • CSV • TXT</div>
    </div>
    """, unsafe_allow_html=True)

    st.markdown(
        '<div class="section-description">'
        'Upload multiple project artifacts to build the project knowledge base.'
        '</div>',
        unsafe_allow_html=True
    )

    uploaded_files = st.file_uploader(
        "Supported formats: PDF, DOCX, CSV, TXT",
        type=["pdf", "docx", "csv", "txt"],
        accept_multiple_files=True,
        label_visibility="visible"
    )

    if uploaded_files:

        st.success(
            f"{len(uploaded_files)} document(s) selected."
        )

        for file in uploaded_files:

            st.write(
                f"**{file.name}** — {file.size / 1024:.1f} KB"
            )

        st.divider()

        if st.button(
            "Process Documents",
            type="primary"
        ):

            # ---------------------------------------------
            # EXTRACT DOCUMENTS
            # ---------------------------------------------

            with st.spinner(
                "Extracting text from documents..."
            ):

                st.session_state.processed_documents = []

                processed_documents = (
                    st.session_state.processed_documents
                )

                for file in uploaded_files:

                    try:

                        text = extract_text(file)

                        processed_documents.append({
                            "filename": file.name,
                            "file_type": file.name.split(".")[-1].upper(),
                            "text": text
                        })

                    except Exception as e:

                        st.error(
                            f"Could not process {file.name}: {e}"
                        )

            # ---------------------------------------------
            # IF PROCESSING SUCCESSFUL
            # ---------------------------------------------

            if processed_documents:

                st.session_state.processed_documents = (
                    processed_documents
                )

                # -----------------------------------------
                # SUCCESS HEADER
                # -----------------------------------------

                st.html(f"""
                <div style="
                    margin-top:25px;
                    padding:24px;
                    border-radius:18px;
                    background:linear-gradient(
                        135deg,
                        #151922,
                        #1B2233
                    );
                    border:1px solid #30384A;
                    text-align:center;
                ">

                    <div style="
                        font-size:30px;
                        margin-bottom:8px;
                    ">
                        ✅
                    </div>

                    <div style="
                        font-size:24px;
                        font-weight:700;
                        color:#F5F5F5;
                    ">
                        Document Processing Complete
                    </div>

                    <div style="
                        margin-top:8px;
                        color:#9AA3B5;
                        font-size:14px;
                    ">
                        {len(processed_documents)}
                        document(s) successfully processed
                    </div>

                </div>
                """)

                # -----------------------------------------
                # CREATE CHUNKS
                # -----------------------------------------

                chunks = split_documents(
                    processed_documents
                )

                st.session_state.chunks = chunks

                # -----------------------------------------
                # GENERATE EMBEDDINGS
                # -----------------------------------------

                with st.spinner(
                    "Generating embeddings..."
                ):

                    embeddings = generate_embeddings(
                        chunks
                    )

                    st.session_state.embeddings = embeddings

                # -----------------------------------------
                # INDEX IN CHROMADB
                # -----------------------------------------

                with st.spinner(
                    "Indexing embeddings in ChromaDB..."
                ):

                    indexed_count = store_embeddings(
                        chunks,
                        embeddings
                    )

                    st.session_state.indexed_count = (
                        indexed_count
                    )

                # -----------------------------------------
                # KNOWLEDGE PIPELINE
                # -----------------------------------------

                st.markdown(
                    "### Knowledge Base Pipeline"
                )

                col1, col2, col3, col4 = st.columns(4)

                with col1:

                    st.html(f"""
                    <div style="
                        background:#151922;
                        border:1px solid #2A303D;
                        border-radius:16px;
                        padding:20px 10px;
                        text-align:center;
                        min-height:145px;
                    ">

                        <div style="font-size:32px;">
                            📄
                        </div>

                        <div style="
                            color:#F5F5F5;
                            font-size:16px;
                            font-weight:650;
                            margin:8px 0;
                        ">
                            Documents
                        </div>

                        <div style="
                            color:#67D5FF;
                            font-size:24px;
                            font-weight:700;
                        ">
                            {len(processed_documents)}
                        </div>

                        <div style="
                            color:#8F98AA;
                            font-size:12px;
                        ">
                            Processed
                        </div>

                    </div>
                    """)

                with col2:

                    st.html(f"""
                    <div style="
                        background:#151922;
                        border:1px solid #2A303D;
                        border-radius:16px;
                        padding:20px 10px;
                        text-align:center;
                        min-height:145px;
                    ">

                        <div style="font-size:32px;">
                            ✂️
                        </div>

                        <div style="
                            color:#F5F5F5;
                            font-size:16px;
                            font-weight:650;
                            margin:8px 0;
                        ">
                            Text Chunks
                        </div>

                        <div style="
                            color:#A78BFA;
                            font-size:24px;
                            font-weight:700;
                        ">
                            {len(chunks)}
                        </div>

                        <div style="
                            color:#8F98AA;
                            font-size:12px;
                        ">
                            Created
                        </div>

                    </div>
                    """)

                with col3:

                    st.html(f"""
                    <div style="
                        background:#151922;
                        border:1px solid #2A303D;
                        border-radius:16px;
                        padding:20px 10px;
                        text-align:center;
                        min-height:145px;
                    ">

                        <div style="font-size:32px;">
                            🧠
                        </div>

                        <div style="
                            color:#F5F5F5;
                            font-size:16px;
                            font-weight:650;
                            margin:8px 0;
                        ">
                            Embeddings
                        </div>

                        <div style="
                            color:#60A5FA;
                            font-size:24px;
                            font-weight:700;
                        ">
                            {len(embeddings)}
                        </div>

                        <div style="
                            color:#8F98AA;
                            font-size:12px;
                        ">
                            Generated
                        </div>

                    </div>
                    """)

                with col4:

                    st.html(f"""
                    <div style="
                        background:#151922;
                        border:1px solid #2A303D;
                        border-radius:16px;
                        padding:20px 10px;
                        text-align:center;
                        min-height:145px;
                    ">

                        <div style="font-size:32px;">
                            🗄️
                        </div>

                        <div style="
                            color:#F5F5F5;
                            font-size:16px;
                            font-weight:650;
                            margin:8px 0;
                        ">
                            ChromaDB
                        </div>

                        <div style="
                            color:#4ADE80;
                            font-size:24px;
                            font-weight:700;
                        ">
                            {indexed_count}
                        </div>

                        <div style="
                            color:#8F98AA;
                            font-size:12px;
                        ">
                            Indexed
                        </div>

                    </div>
                    """)

                # -----------------------------------------
                # FINAL STATUS
                # -----------------------------------------

                st.html("""
                <div style="
                    margin-top:22px;
                    padding:14px;
                    border-radius:12px;
                    background:#13261C;
                    border:1px solid #245C3A;
                    text-align:center;
                    color:#4ADE80;
                    font-weight:600;
                ">
                    ✓ Knowledge Base Ready
                    &nbsp; • &nbsp;
                    Documents are ready for RAG analysis
                </div>
                """)

    # -----------------------------------------------------
    # BACK BUTTON
    # -----------------------------------------------------

    st.divider()

    documents_navigation_columns = st.columns([1, 2, 1])
    with documents_navigation_columns[0]:
        if st.button("← Back to Home"):
            st.session_state.selected_page = "Home"
            st.rerun()
    with documents_navigation_columns[2]:
        if st.button("Go to RAG Pipeline →", key="documents_next_rag", use_container_width=True):
            st.session_state.selected_page = "RAG Pipeline"
            st.rerun()

    st.stop()


# ---------------------------------------------------------
# RAG PIPELINE
# ---------------------------------------------------------

if selected_page == "RAG Pipeline":

    st.markdown("""
    <div class="module-card" style="--accent:#7c9cff;--progress:90%;">
      <div class="module-icon">🔗</div>
      <div class="module-title">RAG Pipeline</div>
      <span class="module-badge">Knowledge Processing</span>

      <div class="module-description">
        Visualize how project documents are transformed into searchable
        knowledge and retrieved for AI-powered project intelligence.
      </div>

      <div class="module-line"><span></span></div>

      <div class="module-caption">
        Upload → Extract → Chunk → Embed → Index → Retrieve
      </div>
    </div>
    """, unsafe_allow_html=True)

    st.markdown(
        '<div class="section-description">'
        'Your project knowledge flows through these stages before being '
        'used by the AI Project Assistant.'
        '</div>',
        unsafe_allow_html=True
    )

    # -----------------------------------------------------
    # PIPELINE STATUS
    # -----------------------------------------------------

    processed_documents = st.session_state.get(
        "processed_documents", []
    )

    chunks = st.session_state.get(
        "chunks", []
    )

    embeddings = st.session_state.get(
        "embeddings"
    )

    indexed_count = st.session_state.get(
        "indexed_count", 0
    )

    upload_status = (
        "Completed"
        if processed_documents
        else "Waiting"
    )

    extract_status = (
        "Completed"
        if processed_documents
        else "Waiting"
    )

    chunk_status = (
        "Completed"
        if chunks
        else "Waiting"
    )

    embed_status = (
        "Completed"
        if embeddings is not None and len(embeddings) > 0
        else "Waiting"
    )

    index_status = (
        "Completed"
        if indexed_count > 0
        else "Waiting"
    )

    pipeline_steps = [
        ("01", "📄", "Upload", upload_status),
        ("02", "📖", "Extract", extract_status),
        ("03", "✂️", "Chunk", chunk_status),
        ("04", "🧠", "Embed", embed_status),
        ("05", "🗄️", "Index", index_status),
    ]

    # -----------------------------------------------------
    # PIPELINE VISUAL
    # -----------------------------------------------------

    st.markdown("### 🔗 Knowledge Processing Pipeline")

    pipeline_cols = st.columns(5)

    for col, (number, icon, name, status) in zip(
        pipeline_cols,
        pipeline_steps
    ):

        with col:

            st.html(
                f"""
                <div style="
                    min-height:155px;
                    padding:20px 12px;
                    border:1px solid #30384A;
                    border-radius:16px;
                    background:linear-gradient(
                        135deg,
                        #151922,
                        #1B2233
                    );
                    text-align:center;
                    box-shadow:0 6px 18px rgba(0,0,0,.15);
                ">

                    <div style="
                        font-size:12px;
                        color:#7d8799;
                        margin-bottom:8px;
                    ">
                        STEP {number}
                    </div>

                    <div style="
                        font-size:30px;
                        margin-bottom:8px;
                    ">
                        {icon}
                    </div>

                    <div style="
                        font-size:17px;
                        font-weight:700;
                        color:#F5F5F5;
                        margin-bottom:12px;
                    ">
                        {name}
                    </div>

                    <div style="
                        display:inline-block;
                        padding:5px 12px;
                        border-radius:20px;
                        font-size:12px;
                        font-weight:600;
                        background:#16351F;
                        color:#4ADE80;
                    ">
                        {status}
                    </div>

                </div>
                """
            )

    # -----------------------------------------------------
    # RAG RETRIEVAL + LLM ANSWER
    # -----------------------------------------------------

    st.markdown("---")

    st.html("""
    <div style="
        padding:22px 24px;
        border:1px solid #38445c;
        border-left:6px solid #8B5CF6;
        border-radius:16px;
        background:linear-gradient(135deg,#171d2b,#20283a);
        margin:20px 0;
    ">

        <div style="
            font-size:28px;
            margin-bottom:8px;
        ">
            🤖
        </div>

        <div style="
            font-size:24px;
            font-weight:750;
            color:#F5F5F5;
            margin-bottom:8px;
        ">
            RAG Query & AI Answer
        </div>

        <div style="
            font-size:14px;
            line-height:1.6;
            color:#9AA3B5;
        ">
            Ask a question about your project documents.
            The system retrieves relevant project information
            and uses it to generate an AI-powered answer.
        </div>

    </div>
    """)

    rag_query = st.text_input(
        "Project Query",
        placeholder="Example: What are the current project blockers?",
        key="rag_pipeline_query"
    )

    if st.button(
        "🔍 Ask AI",
        type="primary",
        use_container_width=True
    ):

        if not processed_documents:

            st.warning(
                "No project documents have been processed yet. "
                "Please upload and process documents first."
            )

        elif not rag_query.strip():

            st.warning(
                "Please enter a project question."
            )

        else:

            # ---------------------------------------------
            # RETRIEVE RELEVANT DOCUMENTS
            # ---------------------------------------------

            with st.spinner(
                "Searching project knowledge base..."
            ):

                results = retrieve_documents(
                    rag_query,
                    top_k=5
                )

            if results:

                st.success(
                    f"{len(results)} relevant project chunks retrieved."
                )

                # -----------------------------------------
                # PREPARE CONTEXT
                # -----------------------------------------

                context = _build_rag_context(results)

                # -----------------------------------------
                # LLM ANSWER
                # -----------------------------------------

                st.markdown("### 🤖 AI Answer")

                with st.spinner(
                    "Generating AI answer..."
                ):

                    answer = generate_answer(
                        rag_query,
                        context
                    )

                st.markdown(
                    f"""
                    <div style="
                        padding:22px;
                        border:1px solid #38445c;
                        border-radius:16px;
                        background:#151922;
                        line-height:1.7;
                        color:#F5F5F5;
                    ">
                        {answer}
                    </div>
                    """,
                    unsafe_allow_html=True
                )

            else:

                st.warning(
                    "No relevant information found in the "
                    "project knowledge base."
                )

    # -----------------------------------------------------
    # HOW RAG WORKS
    # -----------------------------------------------------

    with st.expander(
        "🧠 How RAG Works in This Project"
    ):

        st.markdown(
            "### 📄 1. Document Knowledge"
        )

        st.write(
            "Project documents are extracted and divided "
            "into smaller searchable text chunks."
        )

        st.markdown(
            "### 🧠 2. Semantic Retrieval"
        )

        st.write(
            "When a user asks a question, the system retrieves "
            "the most relevant information from the project "
            "knowledge base."
        )

        st.markdown(
            "### 🤖 3. AI Grounding"
        )

        st.write(
            "The retrieved project context is provided to "
            "the AI model before generating the final answer."
        )

        st.markdown(
            "### 🔄 RAG Flow"
        )

        st.code(
            "Documents → Chunks → Embeddings → ChromaDB "
            "→ Relevant Context → AI Answer"
        )

    # -----------------------------------------------------
    # BACK TO HOME
    # -----------------------------------------------------

    st.markdown("---")

    rag_navigation_columns = st.columns([1, 2, 1])
    with rag_navigation_columns[0]:
        if st.button("← Back to Home"):
            st.session_state.selected_page = "Home"
            st.rerun()
    with rag_navigation_columns[2]:
        if st.button("Go to Project Scope →", key="rag_next_scope", use_container_width=True):
            _prepare_project_scope_entry()
            st.session_state.selected_page = "Project Scope"
            st.rerun()

    st.stop()


# ---------------------------------------------------------
# AI PROJECT INTELLIGENCE ASSISTANT
# ---------------------------------------------------------

if st.session_state.selected_page == "AI Assistant":

    st.markdown(
        """
        <style>
        div[data-testid="stChatMessage"] {
            border: 1px solid rgba(124, 156, 255, .28);
            border-radius: 16px;
            background: linear-gradient(135deg, rgba(23, 29, 43, .94), rgba(22, 32, 52, .88));
            box-shadow: 0 10px 28px rgba(3, 7, 18, .22), inset 0 1px rgba(255,255,255,.04);
        }
        </style>
        """,
        unsafe_allow_html=True
    )

    st.markdown("# AI Project Intelligence Assistant")
    st.caption(
        "Ask questions about your project and get answers grounded in uploaded project documents."
    )

    processed_documents = st.session_state.get("processed_documents", []) or []
    indexed_count = st.session_state.get("indexed_count", 0) or 0

    if not processed_documents or not indexed_count:
        st.info("Upload and process project documents first to use the AI Assistant.")
    else:
        question = st.text_input(
            "Your question",
            placeholder="Ask about project status, risks, blockers, timeline, or actions...",
            key="ai_assistant_question"
        )

        if st.button(
            "Ask",
            type="primary",
            use_container_width=True,
            key="ai_assistant_ask_button"
        ):
            if not question.strip():
                st.warning("Please enter a project question.")
            else:
                with st.spinner("Searching project knowledge base..."):
                    results = retrieve_documents(question, top_k=5)

                if results:
                    context = _build_rag_context(results)

                    with st.spinner("Generating an answer from project context..."):
                        answer = generate_answer(question, context)

                    with st.chat_message("assistant"):
                        st.markdown("**Assistant**")
                        st.markdown(answer)

                    source_names = list(dict.fromkeys(
                        str(result.get("filename", "")).strip()
                        for result in results
                        if result.get("filename")
                    ))
                    if source_names:
                        with st.expander("Retrieved Documents"):
                            for source_name in source_names:
                                st.write(source_name)
                else:
                    st.warning(
                        "No relevant information found in the project knowledge base."
                    )

    st.divider()
    ai_navigation_columns = st.columns([1, 2, 1])
    with ai_navigation_columns[0]:
        if st.button("← Back to Home", key="ai_assistant_back_home"):
            st.session_state.selected_page = "Home"
            st.rerun()
    with ai_navigation_columns[2]:
        if st.button("Go to Documents →", key="ai_assistant_next_documents", use_container_width=True):
            st.session_state.selected_page = "Documents"
            st.rerun()

    st.stop()


# -----------------------------------------------------
# PROJECT SCOPE
# -----------------------------------------------------

if selected_page == "Project Scope":

    st.html("""
    <div class="module-card" style="--accent:#7c9cff;--progress:72%;">

        <div class="module-icon">🎯</div>

        <div class="module-title">
            Project Scope & Deliverables
        </div>

        <span class="module-badge">
            Scope Intelligence
        </span>

        <div class="module-description">
            Analyze project documents to identify the project objective,
            scope, modules, deliverables, requirements, and milestones.
        </div>

        <div class="module-line">
            <span></span>
        </div>

        <div class="module-caption">
            Objective • Modules • Deliverables • Milestones
        </div>

    </div>
    """)

    st.markdown("### 🎯 Analyze Project Scope")

    if st.button(
        "🎯 Analyze Project Scope",
        type="primary",
        use_container_width=True,
        key="analyze_scope_button"
    ):

        if not st.session_state.get("processed_documents"):

            st.warning(
                "No project documents have been processed yet. "
                "Please go to Documents, upload your project files, "
                "and process them first."
            )

        else:

            scope_chunks = retrieve_documents(
                "project objective scope modules features "
                "deliverables requirements deadlines milestones",
                top_k=6
            )
            st.session_state.scope_analysis_attempted = True

            if scope_chunks:

                with st.spinner(
                    "Analyzing project scope..."
                ):

                    scope_analysis = analyze_scope(
                        scope_chunks
                    )

                    st.session_state.scope_analysis = (
                        scope_analysis
                    )

                st.success(
                    "Project scope analysis completed."
                )

            else:

                st.warning(
                    "No relevant project information found "
                    "in the uploaded project documents."
                )

    if st.session_state.get("processed_documents"):
        if not st.session_state.get("scope_analysis_attempted"):
            scope_chunks = retrieve_documents(
                "project objective scope modules features "
                "deliverables requirements deadlines milestones",
                top_k=6
            )
            st.session_state.scope_analysis_attempted = True
            if scope_chunks:
                with st.spinner("Analyzing project scope..."):
                    st.session_state.scope_analysis = analyze_scope(scope_chunks)
            else:
                st.warning("No relevant project information found in the uploaded project documents.")

        if st.session_state.get("scope_analysis"):
            if not st.session_state.get("risk_analysis") and not st.session_state.get("risk_analysis_attempted"):
                risk_chunks = retrieve_documents(
                    "project risks blockers delays dependencies "
                    "schedule delivery challenges issues problems",
                    top_k=6
                )
                st.session_state.risk_analysis_attempted = True
                if risk_chunks:
                    with st.spinner("Analyzing project risks..."):
                        st.session_state.risk_analysis = analyze_risks(risk_chunks)
                else:
                    st.warning("No relevant project risk information found in the uploaded project documents.")

            if not st.session_state.get("blocker_analysis") and not st.session_state.get("blocker_analysis_attempted"):
                blocker_chunks = retrieve_documents(
                    "current blockers unresolved issues pending decisions "
                    "action items dependencies delays problems "
                    "testing integration development schedule",
                    top_k=10
                )
                st.session_state.blocker_analysis_attempted = True
                if blocker_chunks:
                    with st.spinner("Analyzing project blockers..."):
                        st.session_state.blocker_analysis = analyze_blockers(blocker_chunks)
                else:
                    st.warning("No relevant blocker information found in the uploaded project documents.")

    # -------------------------------------------------
    # DISPLAY SCOPE ANALYSIS
    # -------------------------------------------------

    if st.session_state.get("scope_analysis"):

        st.markdown(
            "### 📋 Project Scope Analysis"
        )

        with st.expander(
            "View Complete Scope Analysis",
            expanded=True
        ):
            scope_text = _response_text(st.session_state.scope_analysis)
            scope_text = scope_text.replace("\\n", "\n").replace("<br />", "\n").replace("<br/>", "\n").replace("<br>", "\n")
            scope_section_names = (
                "Project Overview", "Project Objective", "Project Scope",
                "Main Modules / Features", "Key Deliverables",
                "Important Requirements", "Important Deadlines / Milestones",
            )
            scope_name_pattern = "|".join(re.escape(name).replace(r"/", r"\s*/\s*") for name in scope_section_names)
            scope_heading_pattern = re.compile(
                rf"(?im)(?:^|\|)\s*(?:\d+\s*\|\s*)?(?:\*\*)?(?P<title>{scope_name_pattern})(?:\*\*)?\s*\|"
            )
            final_note_match = re.search(
                r"(?is)(If any of the above items were not explicitly mentioned in the documents.*)$",
                scope_text,
            )
            final_note = final_note_match.group(1).strip().strip("|") if final_note_match else ""
            scope_sections_text = scope_text.replace(final_note_match.group(1), "", 1) if final_note_match else scope_text
            scope_matches = list(scope_heading_pattern.finditer(scope_sections_text))
            scope_sections = {}
            for match_index, section_match in enumerate(scope_matches):
                next_start = scope_matches[match_index + 1].start() if match_index + 1 < len(scope_matches) else len(scope_sections_text)
                section_value = scope_sections_text[section_match.end():next_start].strip(" \t\r\n|")
                section_title = re.sub(r"\s*/\s*", " / ", section_match.group("title"))
                section_value = re.sub(r"\s*\|\s*", "\n", section_value).strip()
                if section_value:
                    scope_sections[section_title.casefold()] = section_value

            def scope_list_items(value, split_commas=False, split_sentences=False):
                value = value.replace("**", "").replace("`", "").strip()
                value = re.sub(r"(?im)^\s*(?:[-*•]|\d+[.)])\s*", "", value)
                parts = [part.strip() for part in re.split(r"[\n;•]+", value) if part.strip()]
                if len(parts) == 1 and split_commas:
                    parts = [part.strip() for part in parts[0].split(",") if part.strip()]
                if len(parts) == 1 and split_sentences:
                    parts = [part.strip() for part in re.split(r"(?<=[.!?])\s+(?=[A-Z0-9✓])", parts[0]) if part.strip()]
                return parts

            def scope_card(title):
                st.markdown(f"#### {html.escape(title.upper())}")
                st.markdown("<div style='height:2px;width:48px;margin:-7px 0 12px;background:linear-gradient(90deg,#67d5ff,#b58cff);border-radius:4px;'></div>", unsafe_allow_html=True)

            overview = scope_sections.get("project overview", "")
            objective = scope_sections.get("project objective", "")
            project_name_match = re.search(r"(?i)\bproject\s*(?:name|title)\s*[:：]\s*(.+?)(?=\s*(?:\||;|\n)\s*(?:project\s*)?(?:objective|scope)\s*[:：]|$)", overview)
            objective_match = re.search(r"(?i)\b(?:project\s*)?objective\s*[:：]\s*(.+?)(?=\s*(?:\||;|\n)\s*(?:project\s*)?(?:name|scope)\s*[:：]|$)", overview)
            project_name = project_name_match.group(1).strip(" *|") if project_name_match else ""
            if not objective and objective_match:
                objective = objective_match.group(1).strip(" *|")

            if project_name:
                st.markdown(f"### {html.escape(project_name)}")

            if overview or objective:
                with st.container(border=True):
                    scope_card("Project Overview")
                    if project_name:
                        st.markdown(f"**Project Name:** {project_name}")
                    if objective:
                        st.markdown("**Project Objective**")
                        st.markdown(objective)
                    remaining_overview = overview
                    for matched_value in (project_name_match, objective_match):
                        if matched_value:
                            remaining_overview = remaining_overview.replace(matched_value.group(0), "", 1)
                    remaining_overview = re.sub(r"(?i)\b(?:project\s*)?(?:name|title|objective)\s*[:：]", "", remaining_overview).strip(" *|;\n")
                    if remaining_overview:
                        st.markdown(remaining_overview)

            display_sections = (
                ("project scope", "Project Scope", True, False, False),
                ("main modules / features", "Main Modules / Features", True, False, False),
                ("key deliverables", "Key Deliverables", False, False, False),
                ("important requirements", "Important Requirements", False, True, True),
                ("important deadlines / milestones", "Important Deadlines / Milestones", False, False, False),
            )
            for section_key, section_title, split_commas, checklist, split_sentences in display_sections:
                section_value = scope_sections.get(section_key, "")
                if not section_value:
                    continue
                with st.container(border=True):
                    scope_card(section_title)
                    if section_key == "important deadlines / milestones":
                        sprint_matches = list(re.finditer(r"(?i)(?:^|[;\n])\s*(Sprint\s+\d+)\s*[:：-]?\s*", section_value))
                    else:
                        sprint_matches = []
                    if sprint_matches:
                        for sprint_index, sprint_match in enumerate(sprint_matches):
                            sprint_end = sprint_matches[sprint_index + 1].start() if sprint_index + 1 < len(sprint_matches) else len(section_value)
                            sprint_content = section_value[sprint_match.end():sprint_end].strip(" ;|\n")
                            st.markdown(f"**{sprint_match.group(1).title()}**")
                            for item in scope_list_items(sprint_content):
                                st.markdown(f"- {item}")
                    else:
                        items = scope_list_items(section_value, split_commas=split_commas, split_sentences=split_sentences)
                        for item in items:
                            st.markdown(f"{'✓' if checklist else '-'} {item}")

            if final_note:
                st.info(final_note)
            elif not scope_sections:
                st.markdown(scope_sections_text.replace("|", "\n"))

    else:

        st.info(
            "No scope analysis is available yet. "
            "Open Project Scope with processed documents to generate it."
        )

    if st.session_state.get("risk_analysis"):
        st.markdown("### ⚠️ Project Risks")
        with st.expander("View Complete Risk Analysis", expanded=True):
            st.markdown(_response_text(st.session_state.risk_analysis))
    else:
        st.info("Risk analysis will appear here when relevant project information is available.")

    if st.session_state.get("blocker_analysis"):
        st.markdown("### 🚧 Project Blockers & Actions")
        with st.expander("View Complete Blocker Analysis", expanded=True):
            st.markdown(_response_text(st.session_state.blocker_analysis))
    else:
        st.info("Blocker analysis will appear here when relevant project information is available.")

    # -------------------------------------------------
    # BACK TO HOME
    # -------------------------------------------------

    st.divider()

    scope_navigation_columns = st.columns([1, 2, 1])
    with scope_navigation_columns[0]:
        if st.button(
            "← Back to Home",
            key="scope_back_home"
        ):
            st.session_state.selected_page = "Home"
            st.rerun()
    with scope_navigation_columns[2]:
        if st.button("Go to Risks & Forecast →", key="scope_next_risks", use_container_width=True):
            st.session_state.selected_page = "Risks & Forecast"
            st.rerun()

    st.stop()


# -----------------------------------------------------
# RISKS & DELIVERY FORECAST
# -----------------------------------------------------

if selected_page == "Risks & Forecast":

    st.html("""
    <div class="module-card" style="--accent:#ff9d5c;--progress:65%;">

        <div class="module-icon">⚠️</div>

        <div class="module-title">
            Risks & Delivery Forecast
        </div>

        <span class="module-badge">
            Risk Intelligence
        </span>

        <div class="module-description">
            Identify risks, severity, potential impact,
            dependencies, and delivery challenges from
            the project documents.
        </div>

        <div class="module-line">
            <span></span>
        </div>

        <div class="module-caption">
            Severity • Impact • Dependencies • Delivery
        </div>

    </div>
    """)

    # -------------------------------------------------
    # RISK ANALYSIS RESULT
    # -------------------------------------------------

    if st.session_state.get("risk_analysis"):

        st.markdown("### 📋 Risk Analysis")

        with st.expander(
            "View Complete Risk Analysis",
            expanded=True
        ):
            st.markdown(
                _response_text(st.session_state.risk_analysis)
            )
    else:
        st.info("Open Project Scope to generate the project risk analysis.")

    # -------------------------------------------------
    # BACK TO HOME
    # -------------------------------------------------

    st.divider()

    risk_navigation_columns = st.columns([1, 2, 1])
    with risk_navigation_columns[0]:
        if st.button(
            "← Back to Home",
            key="risk_back_home"
        ):
            st.session_state.selected_page = "Home"
            st.rerun()
    with risk_navigation_columns[2]:
        if st.button("Go to Blockers & Actions →", key="risk_next_blockers", use_container_width=True):
            st.session_state.selected_page = "Blockers & Actions"
            st.rerun()

    st.stop()    


# -----------------------------------------------------
# BLOCKERS & ACTION ITEMS
# -----------------------------------------------------

if selected_page == "Blockers & Actions":

    st.html("""
    <div class="module-card"
         style="--accent:#ff6f91;--progress:65%;">

        <div class="module-icon">🚧</div>

        <div class="module-title">
            Blockers & Action Items
        </div>

        <span class="module-badge">
            Blocker Detection
        </span>

        <div class="module-description">
            Identify current project blockers, their impact,
            required actions, priorities, and suggested owners
            from the project documents.
        </div>

        <div class="module-line">
            <span></span>
        </div>

        <div class="module-caption">
            Issues • Impact • Priorities • Owners • Actions
        </div>

    </div>
    """)

# -------------------------------------------------
# DISPLAY BLOCKERS
# -------------------------------------------------

if selected_page == "Blockers & Actions":

    if st.session_state.get("blocker_analysis"):

        st.markdown(
            "### 📋 Blockers & Recommended Actions"
        )

        blocker_text = str(
            st.session_state.get(
                "blocker_analysis",
                ""
            )
        )

        blocker_text = blocker_text.replace(
            "\\n",
            "\n"
        )

        blocker_text = blocker_text.replace(
            "<br>",
            "\n"
        )

        blocker_text = blocker_text.replace(
            "<br/>",
            "\n"
        )

        blocker_text = blocker_text.replace(
            "<br />",
            "\n"
        )

        blocker_sections = re.split(
            r"(?=###\s+Blocker\s+\d+)",
            blocker_text
        )

        blocker_number = 0

        for section in blocker_sections:

            section = section.strip()

            if not section:
                continue

            blocker_match = re.search(
                r"\*\*Blocker:\*\*\s*(.+)",
                section
            )

            if not blocker_match:
                continue

            description_match = re.search(
                r"\*\*Description:\*\*\s*(.+)",
                section
            )

            priority_match = re.search(
                r"\*\*Priority:\*\*\s*(.+)",
                section
            )

            owner_match = re.search(
                r"\*\*Owner:\*\*\s*(.+)",
                section
            )

            action_match = re.search(
                r"\*\*Required Action:\*\*\s*(.+)",
                section
            )

            status_match = re.search(
                r"\*\*Status:\*\*\s*(.+)",
                section
            )

            blocker_number += 1

            blocker = blocker_match.group(1).strip()

            description = (
                description_match.group(1).strip()
                if description_match
                else "Not specified in the uploaded documents"
            )

            priority = (
                priority_match.group(1).strip()
                if priority_match
                else "Not specified in the uploaded documents"
            )

            owner = (
                owner_match.group(1).strip()
                if owner_match
                else "Not specified in the uploaded documents"
            )

            action = (
                action_match.group(1).strip()
                if action_match
                else "Not specified in the uploaded documents"
            )

            status = (
                status_match.group(1).strip()
                if status_match
                else "Not specified in the uploaded documents"
            )

            priority_lower = priority.lower()

            if "high" in priority_lower:

                border_color = "#ff4d6d"
                priority_icon = "🔴"

            elif "medium" in priority_lower:

                border_color = "#ff9f43"
                priority_icon = "🟠"

            elif "low" in priority_lower:

                border_color = "#2ecc71"
                priority_icon = "🟢"

            else:

                border_color = "#8b5cf6"
                priority_icon = "🔵"

            blocker = html.escape(blocker)
            description = html.escape(description)
            priority = html.escape(priority)
            owner = html.escape(owner)
            action = html.escape(action)
            status = html.escape(status)

            st.html(
                f"""
                <div style="
                    border:1px solid #303746;
                    border-left:6px solid {border_color};
                    border-radius:16px;
                    padding:24px;
                    margin:18px 0;
                    background:linear-gradient(
                        135deg,
                        #151922 0%,
                        #1b2130 100%
                    );
                    box-shadow:0 8px 24px rgba(0,0,0,0.18);
                ">

                    <div style="
                        font-size:22px;
                        font-weight:800;
                        color:#ffffff;
                        margin-bottom:16px;
                    ">
                        🚧 Blocker {blocker_number}
                    </div>

                    <div style="
                        font-size:20px;
                        font-weight:700;
                        color:#ffffff;
                        margin-bottom:18px;
                    ">
                        {blocker}
                    </div>

                    <div style="
                        color:#d7dbe5;
                        font-size:15px;
                        line-height:1.6;
                        margin-bottom:14px;
                    ">
                        <b style="color:#ffffff;">
                            Description:
                        </b>
                        {description}
                    </div>

                    <div style="
                        color:#d7dbe5;
                        font-size:15px;
                        margin-bottom:12px;
                    ">
                        <b style="color:#ffffff;">
                            Priority:
                        </b>
                        {priority_icon} {priority}
                    </div>

                    <div style="
                        color:#d7dbe5;
                        font-size:15px;
                        margin-bottom:12px;
                    ">
                        <b style="color:#ffffff;">
                            Owner:
                        </b>
                        {owner}
                    </div>

                    <div style="
                        color:#d7dbe5;
                        font-size:15px;
                        line-height:1.6;
                        margin-bottom:12px;
                    ">
                        <b style="color:#ffffff;">
                            Required Action:
                        </b>
                        {action}
                    </div>

                    <div style="
                        color:#d7dbe5;
                        font-size:15px;
                    ">
                        <b style="color:#ffffff;">
                            Status:
                        </b>
                        {status}
                    </div>

                </div>
                """
            )

    else:

        st.info(
            "No blocker analysis is available yet. "
            "Open Project Scope to generate it from processed documents."
        )

    st.divider()

    blockers_navigation_columns = st.columns([1, 2, 1])
    with blockers_navigation_columns[0]:
        if st.button(
            "← Back to Home",
            key="blockers_back_home"
        ):
            st.session_state.selected_page = "Home"
            st.rerun()
    with blockers_navigation_columns[2]:
        if st.button("Go to Project Health →", key="blockers_next_health", use_container_width=True):
            st.session_state.selected_page = "Project Health"
            st.rerun()

    st.stop()


# -----------------------------------------------------
# PROJECT DOCUMENTATION
# -----------------------------------------------------

if selected_page == "Documentation":

    st.html("""
    <div class="module-card"
         style="--accent:#b58cff;--progress:85%;">

        <div class="module-icon">📄</div>

        <div class="module-title">
            Project Documentation
        </div>

        <span class="module-badge">
            AI Documentation
        </span>

        <div class="module-description">
            Generate complete structured project documentation
            using uploaded project documents and AI project
            intelligence analysis.
        </div>

        <div class="module-line">
            <span></span>
        </div>

        <div class="module-caption">
            Overview • Scope • Requirements • Risks • Actions
            • Health • Architecture • AI • Summary
        </div>

    </div>
    """)

    st.markdown(
        "### 📝 Generate Project Documentation"
    )

    # -------------------------------------------------
    # GENERATE DOCUMENTATION
    # -------------------------------------------------

    if st.button(
        "📝 Generate Project Documentation",
        type="primary",
        use_container_width=True,
        key="generate_documentation_button"
    ):

        if not st.session_state.get("processed_documents"):

            st.warning(
                "No project documents have been processed yet. "
                "Please go to Documents, upload your project "
                "files, and process them first."
            )

        else:

            # -----------------------------------------
            # RETRIEVE DOCUMENT INFORMATION
            # -----------------------------------------

            documentation_chunks = retrieve_documents(
                "project overview objectives problem statement "
                "scope requirements modules features "
                "deliverables milestones user stories "
                "risks blockers action items dependencies "
                "status progress health architecture "
                "technology RAG AI assistant validation",
                top_k=10
            )

            if documentation_chunks:

                with st.spinner(
                    "Generating complete project documentation..."
                ):

                    # ---------------------------------
                    # GET PREVIOUS AGENT RESULTS
                    # ---------------------------------

                    scope_analysis = st.session_state.get(
                        "scope_analysis",
                        ""
                    )

                    risk_analysis = st.session_state.get(
                        "risk_analysis",
                        ""
                    )

                    blocker_analysis = st.session_state.get(
                        "blocker_analysis",
                        ""
                    )

                    health_analysis = st.session_state.get(
                        "health_analysis",
                        ""
                    )

                    # ---------------------------------
                    # GENERATE COMPLETE DOCUMENTATION
                    # ---------------------------------

                    documentation = generate_project_documentation(
                        documentation_chunks,
                        scope_analysis=scope_analysis,
                        risk_analysis=risk_analysis,
                        blocker_analysis=blocker_analysis,
                        health_analysis=health_analysis
                    )

                    st.session_state.documentation = (
                        documentation
                    )

                # -------------------------------------
                # GENERATION RESULT
                # -------------------------------------

                if documentation and str(
                    documentation
                ).strip():

                    st.success(
                        "Project documentation generated successfully."
                    )

                else:

                    st.error(
                        "Documentation generation returned "
                        "an empty response."
                    )

            else:

                st.warning(
                    "No relevant project information found "
                    "in the uploaded project documents."
                )

    # -------------------------------------------------
    # DISPLAY DOCUMENTATION
    # -------------------------------------------------

    documentation = st.session_state.get(
        "documentation",
        ""
    )

    if documentation:

        clean_documentation = _response_text(documentation)
        clean_documentation = clean_documentation.replace("\\n", "\n")
        clean_documentation = clean_documentation.replace("<br />", "\n").replace("<br/>", "\n").replace("<br>", "\n")
        clean_documentation = re.sub(
            r"(?m)^\s*(?:[-*•]\s*)?\*\*\s*[-*•]\s*([^*]+?)\s*\*\*\s*:\s*",
            r"\1: ",
            clean_documentation,
        )
        clean_documentation = re.sub(
            r"(?m)^\s*(?:[-*•]\s*)?\*\*([^*]+?)\*\*\s*:\s*",
            r"\1: ",
            clean_documentation,
        )
        initial_sections = list(re.finditer(
            r"(?m)^#\s+(\d+)\.\s+(USER STORIES|RISK REGISTER|ACTION ITEMS|PROJECT SUMMARY)\s*$",
            clean_documentation,
        ))
        story_heading = next((match for match in initial_sections if match.group(2) == "USER STORIES"), None)
        if story_heading:
            story_end = next((match.start() for match in initial_sections if match.start() > story_heading.start()), len(clean_documentation))
            story_body = clean_documentation[story_heading.end():story_end]
            has_renderable_stories = (
                re.search(r"(?m)^US-\d+\s*$", story_body)
                and re.search(r"(?m)^User Story:\s*.+$", story_body)
                and re.search(r"(?m)^Related Module:\s*.+$", story_body)
            )
            story_context = "\n".join(
                [
                    str(document.get("text", document.get("content", "")))
                    if isinstance(document, dict) else str(document)
                    for document in (st.session_state.get("processed_documents", []) or [])
                ]
                + [str(st.session_state.get("scope_analysis", ""))]
            )
            story_fallback = _fallback_user_stories(story_context) if not has_renderable_stories else ""
            if story_fallback:
                clean_documentation = (
                    clean_documentation[:story_heading.end()]
                    + "\n\n" + story_fallback + "\n\n"
                    + clean_documentation[story_end:]
                )
        initial_sections = list(re.finditer(
            r"(?m)^#\s+(\d+)\.\s+(USER STORIES|RISK REGISTER|ACTION ITEMS|PROJECT SUMMARY)\s*$",
            clean_documentation,
        ))
        section_bodies = {}
        for section_index, heading_match in enumerate(initial_sections):
            section_end = initial_sections[section_index + 1].start() if section_index + 1 < len(initial_sections) else len(clean_documentation)
            section_bodies[heading_match.group(2)] = clean_documentation[heading_match.end():section_end].strip()
        stories = section_bodies.get("USER STORIES", "")
        if not (
            re.search(r"(?m)^US-\d+\s*$", stories)
            and re.search(r"(?m)^User Story:\s*.+$", stories)
            and re.search(r"(?m)^Related Module:\s*.+$", stories)
        ):
            stories = "Not found in the provided project information."
        risk_records = _format_risk_records(
            st.session_state.get("risk_analysis", ""),
            section_bodies.get("RISK REGISTER", ""),
        )
        actions = _clean_action_items(section_bodies.get("ACTION ITEMS", ""))
        summary = _format_project_summary(section_bodies.get("PROJECT SUMMARY", ""))
        if not summary:
            summary = section_bodies.get("PROJECT SUMMARY", "Not found in the provided project information.")
        cleaned_sections = (
            ("USER STORIES", stories),
            ("RISK REGISTER", risk_records or "Not found in the provided project information."),
            ("ACTION ITEMS", actions or "Not found in the provided project information."),
            ("PROJECT SUMMARY", summary),
        )
        clean_documentation = "\n\n".join(
            f"# {index}. {title}\n\n{content}"
            for index, (title, content) in enumerate(cleaned_sections, start=1)
        )
        st.session_state.documentation = clean_documentation

        st.markdown(
            "### 📋 Generated Project Documentation"
        )

        with st.expander(
            "📄 View Complete Project Documentation",
            expanded=True
        ):
            section_matches = list(re.finditer(
                r"(?m)^#\s+(\d+)\.\s+(USER STORIES|RISK REGISTER|ACTION ITEMS|PROJECT SUMMARY)\s*$",
                clean_documentation,
            ))
            for section_index, heading_match in enumerate(section_matches):
                section_number, section_title = heading_match.groups()
                section_end = section_matches[section_index + 1].start() if section_index + 1 < len(section_matches) else len(clean_documentation)
                section_content = clean_documentation[heading_match.end():section_end].strip()
                st.markdown(f"### {section_number}. {section_title.title()}")

                if section_title == "USER STORIES":
                    records = [part.strip() for part in re.split(r"(?m)(?=^US-\d+\s*$)", section_content) if part.strip()]
                    displayed_stories = 0
                    for record in records:
                        fields = dict(re.findall(r"(?m)^(User Story|Related Module):\s*(.+?)\s*$", record))
                        if not fields:
                            continue
                        displayed_stories += 1
                        with st.container(border=True):
                            story_id = record.splitlines()[0].strip()
                            st.caption(story_id)
                            if fields.get("User Story"):
                                st.markdown(fields["User Story"])
                            if fields.get("Related Module"):
                                st.caption(f"Related Module: {fields['Related Module']}")
                    if not displayed_stories:
                        st.caption("No supported user stories were found in the current project information.")
                elif section_title == "RISK REGISTER":
                    records = [part.strip() for part in re.split(r"(?m)(?=^\s*(?:[-*•]\s*)?Risk(?:\s+\d+)?\s*:)", section_content) if part.strip()]
                    for record in records:
                        lines = record.splitlines()
                        with st.container(border=True):
                            risk_title = re.sub(r"^\s*(?:[-*•]\s*)?", "", lines[0].strip())
                            risk_title = risk_title.replace("**", "")
                            st.markdown(f"**{risk_title}**")
                            for line in lines[1:]:
                                match = re.match(r"^\s*([^:]+):\s*(.+?)\s*$", line)
                                if match:
                                    st.markdown(f"**{match.group(1)}:** {match.group(2)}")
                elif section_title == "ACTION ITEMS":
                    records = [part.strip() for part in re.split(r"(?m)(?=^Action\s+\d+\s*:)", section_content) if part.strip()]
                    for record in records:
                        lines = record.splitlines()
                        with st.container(border=True):
                            st.markdown(f"**{lines[0].strip()}**")
                            for line in lines[1:]:
                                match = re.match(r"^\s*([^:]+):\s*(.+?)\s*$", line)
                                if match:
                                    st.caption(f"{match.group(1)}: {match.group(2)}")
                else:
                    for field_match in re.finditer(
                        r"(?ms)^([^:\n]+):\s*\n((?:(?:\s*[-*•]\s+.+|\s+.+)\n?)*)",
                        section_content,
                    ):
                        field_name = field_match.group(1).strip()
                        field_value = field_match.group(2).strip()
                        if field_value:
                            with st.container(border=True):
                                st.markdown(f"**{field_name}**")
                                st.markdown(field_value)

        # -----------------------------------------
        # DOWNLOAD GENERATED DOCUMENTATION
        # -----------------------------------------

        from io import BytesIO
        from docx import Document

        download_text = clean_documentation

        section_heading_pattern = re.compile(
            r"^\s*#*\s*(?:\d+\.\s*)?(USER STORIES|RISK REGISTER|"
            r"ACTION ITEMS|PROJECT SUMMARY)\s*#*\s*$",
            re.IGNORECASE
        )

        docx_document = Document()
        docx_document.add_heading("Project Documentation", 0)

        for line in download_text.splitlines():
            stripped_line = line.strip()
            heading_match = section_heading_pattern.match(stripped_line)

            if heading_match:
                docx_document.add_heading(
                    heading_match.group(1).upper(),
                    level=1
                )
            elif stripped_line.startswith(("- ", "* ")):
                docx_document.add_paragraph(
                    stripped_line[2:],
                    style="List Bullet"
                )
            elif stripped_line:
                docx_document.add_paragraph(
                    stripped_line.replace("**", "").replace("`", "")
                )

        docx_buffer = BytesIO()
        docx_document.save(docx_buffer)
        docx_buffer.seek(0)

        # Build a simple text PDF using only the Python standard library.
        pdf_lines = [("Project Documentation", "title")]

        for line in download_text.splitlines():
            stripped_line = line.strip()
            heading_match = section_heading_pattern.match(stripped_line)

            if heading_match:
                pdf_lines.append((heading_match.group(1).upper(), "heading"))
            elif stripped_line:
                plain_line = stripped_line.replace("**", "").replace("`", "")
                plain_line = plain_line.replace("•", "-").replace("–", "-").replace("—", "-")
                plain_line = plain_line.replace("’", "'").replace("‘", "'")
                plain_line = plain_line.replace("“", '"').replace("”", '"')
                plain_line = plain_line.encode("ascii", "ignore").decode("ascii")

                while len(plain_line) > 92:
                    split_at = plain_line.rfind(" ", 0, 92)
                    if split_at < 1:
                        split_at = 92
                    pdf_lines.append((plain_line[:split_at], "body"))
                    plain_line = plain_line[split_at:].lstrip()

                pdf_lines.append((plain_line, "body"))

        pdf_pages = []
        page_commands = []
        y_position = 790

        for line_text, line_style in pdf_lines:
            if y_position < 55:
                pdf_pages.append("\n".join(page_commands).encode("ascii"))
                page_commands = []
                y_position = 790

            font_name, font_size, line_gap = {
                "title": ("F2", 20, 34),
                "heading": ("F2", 13, 24),
                "body": ("F1", 10, 15)
            }[line_style]
            escaped_text = line_text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            page_commands.append(
                f"BT /{font_name} {font_size} Tf 50 {y_position} Td "
                f"({escaped_text}) Tj ET"
            )
            y_position -= line_gap

        if page_commands:
            pdf_pages.append("\n".join(page_commands).encode("ascii"))

        pdf_objects = [
            b"<< /Type /Catalog /Pages 2 0 R >>",
            b"",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold >>"
        ]
        page_ids = []

        for page_index, page_content in enumerate(pdf_pages):
            page_id = 5 + page_index * 2
            stream_id = page_id + 1
            page_ids.append(page_id)
            pdf_objects.append(
                f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 842] "
                f"/Resources << /Font << /F1 3 0 R /F2 4 0 R >> >> "
                f"/Contents {stream_id} 0 R >>".encode("ascii")
            )
            pdf_objects.append(
                f"<< /Length {len(page_content)} >>\nstream\n".encode("ascii")
                + page_content
                + b"\nendstream"
            )

        pdf_objects[1] = (
            f"<< /Type /Pages /Kids [{ ' '.join(f'{page_id} 0 R' for page_id in page_ids) }] "
            f"/Count {len(page_ids)} >>"
        ).encode("ascii")

        pdf_buffer = bytearray(b"%PDF-1.4\n")
        object_offsets = [0]

        for object_id, pdf_object in enumerate(pdf_objects, start=1):
            object_offsets.append(len(pdf_buffer))
            pdf_buffer.extend(f"{object_id} 0 obj\n".encode("ascii"))
            pdf_buffer.extend(pdf_object)
            pdf_buffer.extend(b"\nendobj\n")

        xref_offset = len(pdf_buffer)
        pdf_buffer.extend(f"xref\n0 {len(object_offsets)}\n".encode("ascii"))
        pdf_buffer.extend(b"0000000000 65535 f \n")
        for object_offset in object_offsets[1:]:
            pdf_buffer.extend(f"{object_offset:010d} 00000 n \n".encode("ascii"))
        pdf_buffer.extend(
            f"trailer\n<< /Size {len(object_offsets)} /Root 1 0 R >>\n"
            f"startxref\n{xref_offset}\n%%EOF".encode("ascii")
        )

        pdf_column, docx_column = st.columns(2)
        with pdf_column:
            st.download_button(
                "Download as PDF",
                data=bytes(pdf_buffer),
                file_name="project_documentation.pdf",
                mime="application/pdf",
                key="documentation_download_pdf"
            )
        with docx_column:
            st.download_button(
                "Download as DOCX",
                data=docx_buffer.getvalue(),
                file_name="project_documentation.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                key="documentation_download_docx"
            )

    # -------------------------------------------------
    # BACK TO HOME
    # -------------------------------------------------

    st.divider()

    documentation_navigation_columns = st.columns([1, 2, 1])
    with documentation_navigation_columns[0]:
        if st.button(
            "← Back to Home",
            key="documentation_back_home"
        ):
            st.session_state.selected_page = "Home"
            st.rerun()
    with documentation_navigation_columns[2]:
        if st.button("Go to Dashboard →", key="documentation_next_dashboard", use_container_width=True):
            st.session_state.selected_page = "Dashboard"
            st.rerun()

    st.stop()


